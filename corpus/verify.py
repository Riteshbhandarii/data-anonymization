#!/usr/bin/env python3
"""Check that a generated corpus matches its labels.

The labels are the whole value of this corpus. If one says an IBAN is in the
hidden sheet and it is not, every number measured against it is wrong and
nothing complains. This reads each document back and confirms every labelled
value really is where the label says.

    python3 corpus/verify.py ./corpus-out
"""

import json
import os
import sys
import zipfile


def text_at(path, location):
    """The text of one part of a document, without a full extraction layer."""
    ext = os.path.splitext(path)[1].lower()
    if location == "metadata":
        if ext == ".pdf":
            from pypdf import PdfReader
            return " ".join(str(v) for v in (PdfReader(path).metadata or {}).values())
        with zipfile.ZipFile(path) as z:
            return " ".join(z.read(n).decode("utf-8", "replace")
                            for n in z.namelist() if n.startswith("docProps/"))
    if location == "notes":
        with zipfile.ZipFile(path) as z:
            return " ".join(z.read(n).decode("utf-8", "replace")
                            for n in z.namelist() if n.startswith("ppt/notesSlides/notesSlide"))
    if location == "hidden_sheet":
        import openpyxl
        wb = openpyxl.load_workbook(path)
        return " ".join(str(c) for ws in wb.worksheets if ws.sheet_state != "visible"
                        for row in ws.iter_rows(values_only=True) for c in row if c is not None)
    if ext == ".docx":
        import docx
        return "\n".join(p.text for p in docx.Document(path).paragraphs)
    if ext == ".pptx":
        from pptx import Presentation
        return "\n".join(s.text_frame.text for sl in Presentation(path).slides
                         for s in sl.shapes if s.has_text_frame)
    if ext == ".xlsx":
        import openpyxl
        wb = openpyxl.load_workbook(path, data_only=True)
        return " ".join(str(c) for ws in wb.worksheets if ws.sheet_state == "visible"
                        for row in ws.iter_rows(values_only=True) for c in row if c is not None)
    if ext == ".pdf":
        from pypdf import PdfReader
        return "\n".join(pg.extract_text() or "" for pg in PdfReader(path).pages)
    return open(path, encoding="utf-8", errors="replace").read()


def main(root):
    labels = os.path.join(root, "labels")
    if not os.path.isdir(labels):
        sys.exit(f"no labels directory in {root}")

    checked = 0
    bad = []
    for name in sorted(os.listdir(labels)):
        with open(os.path.join(labels, name), encoding="utf-8") as f:
            doc = json.load(f)
        path = os.path.join(root, doc["file"])
        if not os.path.exists(path):
            bad.append(f"{doc['file']}: document missing")
            continue
        cache = {}
        for e in doc["entities"]:
            loc = e["location"]
            if loc not in cache:
                cache[loc] = text_at(path, loc)
            checked += 1
            if e["value"] not in cache[loc]:
                bad.append(f"{doc['file']}: {e['type']} not found in {loc}: {e['value'][:40]}")

    print(f"checked {checked} labels across {len(os.listdir(labels))} documents")
    if bad:
        print(f"\n{len(bad)} mismatches:")
        for b in bad[:20]:
            print(f"  {b}")
        sys.exit(1)
    print("every labelled value is where its label says")


if __name__ == "__main__":
    main(sys.argv[1] if len(sys.argv) > 1 else "./corpus-out")
