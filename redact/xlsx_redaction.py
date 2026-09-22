from pathlib import Path
from openpyxl import load_workbook


SENSITIVE_FIELDS = {
    "name",
    "email",
    "phone",
    "company",
    "iban",
    "date",
    "personal_id",
    "plate",
    "address",
}


def redact_xlsx(input_path, output_path):
    input_path = Path(input_path)
    output_path = Path(output_path)

    workbook = load_workbook(input_path)

    # Redact metadata
    workbook.properties.creator = "[REDACTED]"
    workbook.properties.lastModifiedBy = "[REDACTED]"

    # Redact sensitive cells
    for sheet in workbook.worksheets:

        headers = {}

        for cell in sheet[1]:
            if cell.value is not None:
                headers[cell.column] = str(cell.value).strip().lower()

        for row in sheet.iter_rows(min_row=2):

            for cell in row:

                if cell.value is None:
                    continue

                header = headers.get(cell.column)

                if header in SENSITIVE_FIELDS:
                    cell.value = "[REDACTED]"

    output_path.parent.mkdir(
        parents=True,
        exist_ok=True
    )

    workbook.save(output_path)

    print("\nRedacted file saved to:")
    print(output_path)


if __name__ == "__main__":

    input_path = input(
        "Enter original XLSX file path: "
    ).strip().strip('"')

    output_path = input(
        "Enter output XLSX file path: "
    ).strip().strip('"')

    redact_xlsx(input_path, output_path)