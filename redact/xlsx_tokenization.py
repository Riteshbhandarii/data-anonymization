from pathlib import Path
from openpyxl import load_workbook


FIELD_TYPES = {
    "name": "PERSON",
    "email": "EMAIL",
    "phone": "PHONE",
    "company": "COMPANY",
    "iban": "IBAN",
    "date": "DATE",
    "personal_id": "PERSONAL_ID",
    "plate": "PLATE",
    "address": "ADDRESS",
}


def tokenize_xlsx(input_path, output_path):
    input_path = Path(input_path)
    output_path = Path(output_path)

    workbook = load_workbook(input_path)

    # Store original value -> token
    mappings = {}

    # Counter for each data type
    counters = {}

    def get_token(data_type, original_value):
        value = str(original_value)

        key = (data_type, value)

        # Same original value = same token
        if key in mappings:
            return mappings[key]

        counters[data_type] = counters.get(data_type, 0) + 1

        token = f"[{data_type}_{counters[data_type]:03d}]"

        mappings[key] = token

        return token

    # -----------------------------
    # Metadata
    # -----------------------------

    if workbook.properties.creator:
        workbook.properties.creator = get_token(
            "PERSON",
            workbook.properties.creator
        )

    if workbook.properties.lastModifiedBy:
        workbook.properties.lastModifiedBy = get_token(
            "PERSON",
            workbook.properties.lastModifiedBy
        )

    # -----------------------------
    # Worksheet cells
    # -----------------------------

    for sheet in workbook.worksheets:

        headers = {}

        for cell in sheet[1]:
            if cell.value is not None:
                headers[cell.column] = (
                    str(cell.value)
                    .strip()
                    .lower()
                )

        for row in sheet.iter_rows(min_row=2):

            for cell in row:

                if cell.value is None:
                    continue

                header = headers.get(cell.column)

                if header in FIELD_TYPES:

                    data_type = FIELD_TYPES[header]

                    cell.value = get_token(
                        data_type,
                        cell.value
                    )

    # -----------------------------
    # Save output
    # -----------------------------

    output_path.parent.mkdir(
        parents=True,
        exist_ok=True
    )

    workbook.save(output_path)

    print("\nTokenized file saved to:")
    print(output_path)

    print("\n[TOKEN MAPPINGS]")

    for (data_type, original_value), token in mappings.items():
        print(
            f"{data_type:<12} | "
            f"{original_value} -> {token}"
        )


if __name__ == "__main__":

    input_path = input(
        "Enter original XLSX file path: "
    ).strip().strip('"')

    output_path = input(
        "Enter output XLSX file path: "
    ).strip().strip('"')

    tokenize_xlsx(input_path, output_path)