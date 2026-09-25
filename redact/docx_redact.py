import docx
from presidio_analyzer import AnalyzerEngine
from presidio_anonymizer import AnonymizerEngine
from presidio_anonymizer.entities import OperatorConfig
from presidio_anonymizer.operators import Operator, OperatorType
from typing import Dict


class ConsistentMappingOperator(Operator):
    def operator_name(self) -> str: return "ConsistentMappingOperator"
    def operator_type(self) -> OperatorType: return OperatorType.Anonymize

    def operate(self, text: str, params: Dict = None) -> str:
        entity_type = params["entity_type"]
        entity_mapping = params["entity_mapping"]
        if entity_type not in entity_mapping:
            entity_mapping[entity_type] = {}
        mapping = entity_mapping[entity_type]
        if text in mapping: return mapping[text]
        new_placeholder = f"[{entity_type}_{len(mapping) + 1:03d}]"
        mapping[text] = new_placeholder
        return new_placeholder
    def validate(self, params: Dict = None): pass

def redact_docx(input_path, output_path):
    doc = docx.Document(input_path)
    full_text = "\n".join([p.text for p in doc.paragraphs if p.text.strip()])

    analyzer = AnalyzerEngine()
    anonymizer = AnonymizerEngine()
    anonymizer.add_anonymizer(ConsistentMappingOperator)

    detected_pii = analyzer.analyze(text=full_text, language='en')
    translation_key ={}

    sanitized = anonymizer.anonymize(
        text=full_text,
        analyzer_results=detected_pii,
        operators={"DEFUALT": OperatorConfig("ConsistentMappingOperator", {"entity_mapping": translation_key})}
    )

    new_doc = docx.Document()
    for line in sanitized.text.split('\n'):
        new_doc.add_paragraph(line)

    new_doc.save(output_path)
    print(f"\n[SUCCESS] Anonymized document saved to: {output_path}")

if __name__ == "__main__":
    in_path = input("Enter original DOCX path: ").strip().strip('"')
    out_path = input("Enter output DOCX path: ").strip().strip('"')
    redact_docx(in_path, out_path)