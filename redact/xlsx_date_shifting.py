from pathlib import Path
from datetime import datetime, timedelta, date
from openpyxl import load_workbook


FIELD_TYPES = {
    "name": "PERSON",
    "email": "EMAIL",
    "phone": "PHONE",
    "company": "COMPANY",
    "iban": "IBAN",
    "date": "DATE",
    "personal_id": "PERSONAL_ID",
    "plate": "PLATE",
    "address": "ADDRESS",
}


SHIFT_DAYS = 30


def shift_date(value):
    """
    Shift a date by a fixed number of days.
    """

    # Excel date object
    if isinstance(value, (datetime, date)):
        return value + timedelta(days=SHIFT_DAYS)

    # String date: YYYY-MM-DD
    try:
        parsed = datetime.strptime(
            str(value),
            "%Y-%m-%d"
        )

        shifted = parsed + timedelta(days=SHIFT_DAYS)

        return shifted.date()

    except ValueError:
        return "[INVALID DATE]"


def date_shift_xlsx(input_path, output_path):

    input_path = Path(input_path)
    output_path = Path(output_path)

    workbook = load_workbook(input_path)

    mappings = {}

    # Metadata contains personal names,
    # so redact them.
    if workbook.properties.creator:

        original = str(
            workbook.properties.creator
        )

        workbook.properties.creator = "[REDACTED]"

        mappings[
            ("PERSON", original)
        ] = "[REDACTED]"

    if workbook.properties.lastModifiedBy:

        original = str(
            workbook.properties.lastModifiedBy
        )

        workbook.properties.lastModifiedBy = "[REDACTED]"

        mappings[
            ("PERSON", original)
        ] = "[REDACTED]"

    # Worksheets
    for sheet in workbook.worksheets:

        headers = {}

        for cell in sheet[1]:

            if cell.value is not None:

                headers[cell.column] = (
                    str(cell.value)
                    .strip()
                    .lower()
                )

        for row in sheet.iter_rows(min_row=2):

            for cell in row:

                if cell.value is None:
                    continue

                header = headers.get(
                    cell.column
                )

                if header not in FIELD_TYPES:
                    continue

                data_type = FIELD_TYPES[header]

                original = cell.value

                # Date -> shift
                if data_type == "DATE":

                    shifted = shift_date(original)

                    mappings[
                        (
                            data_type,
                            str(original)
                        )
                    ] = str(shifted)

                    cell.value = shifted

                # Other sensitive fields -> redact
                else:

                    mappings[
                        (
                            data_type,
                            str(original)
                        )
                    ] = "[REDACTED]"

                    cell.value = "[REDACTED]"

    output_path.parent.mkdir(
        parents=True,
        exist_ok=True
    )

    workbook.save(output_path)

    print("\nDate-shifted file saved to:")
    print(output_path)

    print(
        f"\nDate shift applied: +{SHIFT_DAYS} days"
    )

    print("\n[DATE SHIFTING MAPPINGS]")

    for (
        data_type,
        original
    ), new_value in mappings.items():

        print(
            f"{data_type:<12} | "
            f"{original} -> {new_value}"
        )


if __name__ == "__main__":

    input_path = input(
        "Enter original XLSX file path: "
    ).strip().strip('"')

    output_path = input(
        "Enter output XLSX file path: "
    ).strip().strip('"')

    date_shift_xlsx(
        input_path,
        output_path
    )