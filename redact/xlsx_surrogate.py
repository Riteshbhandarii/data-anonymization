from pathlib import Path
from openpyxl import load_workbook
from datetime import datetime, date


FIELD_TYPES = {
    "name": "PERSON",
    "email": "EMAIL",
    "phone": "PHONE",
    "company": "COMPANY",
    "iban": "IBAN",
    "date": "DATE",
    "personal_id": "PERSONAL_ID",
    "plate": "PLATE",
    "address": "ADDRESS",
}


# Synthetic replacement values
SURROGATES = {
    "PERSON": [
        "Alex Morgan",
        "Jamie Taylor",
        "Jordan Parker",
        "Casey Wilson",
        "Taylor Brooks",
    ],

    "EMAIL": [
        "alex.morgan@example.com",
        "jamie.taylor@example.com",
        "jordan.parker@example.com",
        "casey.wilson@example.com",
        "taylor.brooks@example.com",
    ],

    "PHONE": [
        "+358 40 000 1001",
        "+358 40 000 1002",
        "+358 40 000 1003",
        "+358 40 000 1004",
        "+358 40 000 1005",
    ],

    "COMPANY": [
        "Northstar Example Ltd.",
        "Bluewave Example Oy",
        "Demo Analytics Ltd.",
        "Example Health Oy",
        "Testworks Ltd.",
    ],

    "IBAN": [
        "FI00TEST000000000001",
        "FI00TEST000000000002",
        "FI00TEST000000000003",
        "FI00TEST000000000004",
        "FI00TEST000000000005",
    ],

    "DATE": [
        "1990-03-12",
        "1995-07-24",
        "2000-11-08",
        "1988-05-18",
        "2002-09-30",
    ],

    "PERSONAL_ID": [
        "TEST-ID-001",
        "TEST-ID-002",
        "TEST-ID-003",
        "TEST-ID-004",
        "TEST-ID-005",
    ],

    "PLATE": [
        "ZZZ-001",
        "ZZZ-002",
        "ZZZ-003",
        "ZZZ-004",
        "ZZZ-005",
    ],

    "ADDRESS": [
        "100 Example Street, Test City",
        "200 Example Avenue, Demo City",
        "300 Sample Road, Test Town",
        "400 Demo Street, Example City",
        "500 Test Avenue, Sample Town",
    ],
}


def surrogate_xlsx(input_path, output_path):

    input_path = Path(input_path)
    output_path = Path(output_path)

    workbook = load_workbook(input_path)

    # original value -> surrogate value
    mappings = {}

    # counter for each data type
    counters = {}

    def get_surrogate(data_type, original_value):

        value = str(original_value)

        key = (data_type, value)

        # Same original value = same surrogate
        if key in mappings:
            return mappings[key]

        index = counters.get(data_type, 0)

        replacements = SURROGATES[data_type]

        # If more values than prepared surrogates,
        # create a simple synthetic fallback
        if index < len(replacements):
            surrogate = replacements[index]
        else:
            surrogate = f"SYNTHETIC_{data_type}_{index + 1:03d}"

        counters[data_type] = index + 1
        mappings[key] = surrogate

        return surrogate

    # --------------------------------------------------
    # 1. Replace metadata
    # --------------------------------------------------

    if workbook.properties.creator:

        workbook.properties.creator = get_surrogate(
            "PERSON",
            workbook.properties.creator
        )

    if workbook.properties.lastModifiedBy:

        workbook.properties.lastModifiedBy = get_surrogate(
            "PERSON",
            workbook.properties.lastModifiedBy
        )

    # --------------------------------------------------
    # 2. Replace worksheet values
    # --------------------------------------------------

    for sheet in workbook.worksheets:

        headers = {}

        for cell in sheet[1]:

            if cell.value is not None:

                header = str(
                    cell.value
                ).strip().lower()

                headers[cell.column] = header

        for row in sheet.iter_rows(min_row=2):

            for cell in row:

                if cell.value is None:
                    continue

                header = headers.get(cell.column)

                if header not in FIELD_TYPES:
                    continue

                data_type = FIELD_TYPES[header]

                surrogate = get_surrogate(
                    data_type,
                    cell.value
                )

                # Keep Excel date cells as actual dates
                if data_type == "DATE":

                    cell.value = datetime.strptime(
                        surrogate,
                        "%Y-%m-%d"
                    ).date()

                else:
                    cell.value = surrogate

    # --------------------------------------------------
    # 3. Save file
    # --------------------------------------------------

    output_path.parent.mkdir(
        parents=True,
        exist_ok=True
    )

    workbook.save(output_path)

    print("\nSurrogate-substituted file saved to:")
    print(output_path)

    # --------------------------------------------------
    # 4. Print mappings
    # --------------------------------------------------

    print("\n[SURROGATE MAPPINGS]")

    for (data_type, original_value), surrogate in mappings.items():

        print(
            f"{data_type:<12} | "
            f"{original_value} -> {surrogate}"
        )


if __name__ == "__main__":

    input_path = input(
        "Enter original XLSX file path: "
    ).strip().strip('"')

    output_path = input(
        "Enter output XLSX file path: "
    ).strip().strip('"')

    surrogate_xlsx(
        input_path,
        output_path
    )