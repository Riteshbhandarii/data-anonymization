from pathlib import Path
from openpyxl import load_workbook


FIELD_TYPES = {
    "name": "PERSON",
    "email": "EMAIL",
    "phone": "PHONE",
    "company": "COMPANY",
    "iban": "IBAN",
    "date": "DATE",
    "personal_id": "PERSONAL_ID",
    "plate": "PLATE",
    "address": "ADDRESS",
}


def mask_value(data_type, value):
    value = str(value)

    if data_type == "PERSON":
        parts = value.split()
        return " ".join(
            part[0] + "*" * max(len(part) - 1, 0)
            for part in parts
        )

    if data_type == "EMAIL":
        if "@" not in value:
            return "*" * len(value)

        username, domain = value.split("@", 1)

        if len(username) <= 1:
            masked_user = "*"
        else:
            masked_user = username[0] + "*" * (len(username) - 1)

        return f"{masked_user}@{domain}"

    if data_type == "PHONE":
        if len(value) <= 4:
            return "*" * len(value)

        return "*" * (len(value) - 4) + value[-4:]

    if data_type == "COMPANY":
        parts = value.split()

        return " ".join(
            part[0] + "*" * max(len(part) - 1, 0)
            for part in parts
        )

    if data_type == "IBAN":
        if len(value) <= 4:
            return "*" * len(value)

        return "*" * (len(value) - 4) + value[-4:]

    if data_type == "DATE":
        return value[:4] + "-**-**"

    if data_type == "PERSONAL_ID":
        if len(value) <= 4:
            return "*" * len(value)

        return "*" * (len(value) - 4) + value[-4:]

    if data_type == "PLATE":
        if len(value) <= 2:
            return "*" * len(value)

        return "*" * (len(value) - 2) + value[-2:]

    if data_type == "ADDRESS":
        return "[MASKED ADDRESS]"

    return "[MASKED]"


def mask_xlsx(input_path, output_path):
    input_path = Path(input_path)
    output_path = Path(output_path)

    workbook = load_workbook(input_path)

    mappings = {}

    # Metadata
    if workbook.properties.creator:
        original = str(workbook.properties.creator)
        masked = mask_value("PERSON", original)
        workbook.properties.creator = masked
        mappings[("PERSON", original)] = masked

    if workbook.properties.lastModifiedBy:
        original = str(workbook.properties.lastModifiedBy)
        masked = mask_value("PERSON", original)
        workbook.properties.lastModifiedBy = masked
        mappings[("PERSON", original)] = masked

    # Cells
    for sheet in workbook.worksheets:

        headers = {}

        for cell in sheet[1]:
            if cell.value is not None:
                headers[cell.column] = str(cell.value).strip().lower()

        for row in sheet.iter_rows(min_row=2):

            for cell in row:

                if cell.value is None:
                    continue

                header = headers.get(cell.column)

                if header not in FIELD_TYPES:
                    continue

                data_type = FIELD_TYPES[header]
                original = str(cell.value)

                key = (data_type, original)

                if key in mappings:
                    masked = mappings[key]
                else:
                    masked = mask_value(data_type, original)
                    mappings[key] = masked

                cell.value = masked

    output_path.parent.mkdir(
        parents=True,
        exist_ok=True
    )

    workbook.save(output_path)

    print("\nMasked file saved to:")
    print(output_path)

    print("\n[MASKING MAPPINGS]")

    for (data_type, original), masked in mappings.items():
        print(
            f"{data_type:<12} | "
            f"{original} -> {masked}"
        )


if __name__ == "__main__":

    input_path = input(
        "Enter original XLSX file path: "
    ).strip().strip('"')

    output_path = input(
        "Enter output XLSX file path: "
    ).strip().strip('"')

    mask_xlsx(
        input_path,
        output_path
    )