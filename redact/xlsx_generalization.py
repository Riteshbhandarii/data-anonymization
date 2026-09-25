from pathlib import Path
from openpyxl import load_workbook
import re


FIELD_TYPES = {
    "name": "PERSON",
    "email": "EMAIL",
    "phone": "PHONE",
    "company": "COMPANY",
    "iban": "IBAN",
    "date": "DATE",
    "personal_id": "PERSONAL_ID",
    "plate": "PLATE",
    "address": "ADDRESS",
}


def generalize_value(data_type, value):
    value = str(value)

    # Keep only entity category for direct personal names
    if data_type == "PERSON":
        return "[PERSON]"

    # Keep only email domain
    if data_type == "EMAIL":
        if "@" in value:
            return value.split("@", 1)[1]

        return "[EMAIL]"

    # Phone number: remove exact number
    if data_type == "PHONE":
        return "[PHONE]"

    # Company: remove exact company name
    if data_type == "COMPANY":
        return "[COMPANY]"

    # Keep IBAN country code only
    if data_type == "IBAN":
        if len(value) >= 2:
            return value[:2] + "-IBAN"

        return "[IBAN]"

    # Keep year only
    if data_type == "DATE":
        return value[:4]

    # Personal identity code:
    # keep birth year information where possible
    if data_type == "PERSONAL_ID":

        match = re.match(
            r"(\d{2})(\d{2})(\d{2})([-+A])",
            value
        )

        if match:
            yy = match.group(3)
            century = match.group(4)

            if century == "+":
                year = "18" + yy
            elif century == "-":
                year = "19" + yy
            elif century == "A":
                year = "20" + yy
            else:
                year = yy

            return f"BIRTH_YEAR_{year}"

        return "[PERSONAL_ID]"

    # Keep only plate category
    if data_type == "PLATE":
        return "[PLATE]"

    # Try to keep US state abbreviation
    if data_type == "ADDRESS":

        match = re.search(
            r",\s*([A-Z]{2})\s+\d{5}",
            value
        )

        if match:
            return f"STATE_{match.group(1)}"

        return "[GENERALIZED_ADDRESS]"

    return "[GENERALIZED]"


def generalize_xlsx(input_path, output_path):

    input_path = Path(input_path)
    output_path = Path(output_path)

    workbook = load_workbook(input_path)

    mappings = {}

    # Metadata
    if workbook.properties.creator:

        original = str(workbook.properties.creator)

        generalized = generalize_value(
            "PERSON",
            original
        )

        workbook.properties.creator = generalized

        mappings[
            ("PERSON", original)
        ] = generalized

    if workbook.properties.lastModifiedBy:

        original = str(
            workbook.properties.lastModifiedBy
        )

        generalized = generalize_value(
            "PERSON",
            original
        )

        workbook.properties.lastModifiedBy = generalized

        mappings[
            ("PERSON", original)
        ] = generalized

    # Worksheets
    for sheet in workbook.worksheets:

        headers = {}

        for cell in sheet[1]:

            if cell.value is not None:

                headers[cell.column] = (
                    str(cell.value)
                    .strip()
                    .lower()
                )

        for row in sheet.iter_rows(min_row=2):

            for cell in row:

                if cell.value is None:
                    continue

                header = headers.get(
                    cell.column
                )

                if header not in FIELD_TYPES:
                    continue

                data_type = FIELD_TYPES[header]
                original = str(cell.value)

                key = (
                    data_type,
                    original
                )

                if key in mappings:

                    generalized = mappings[key]

                else:

                    generalized = generalize_value(
                        data_type,
                        original
                    )

                    mappings[key] = generalized

                cell.value = generalized

    output_path.parent.mkdir(
        parents=True,
        exist_ok=True
    )

    workbook.save(output_path)

    print("\nGeneralized file saved to:")
    print(output_path)

    print("\n[GENERALIZATION MAPPINGS]")

    for (
        data_type,
        original
    ), generalized in mappings.items():

        print(
            f"{data_type:<12} | "
            f"{original} -> {generalized}"
        )


if __name__ == "__main__":

    input_path = input(
        "Enter original XLSX file path: "
    ).strip().strip('"')

    output_path = input(
        "Enter output XLSX file path: "
    ).strip().strip('"')

    generalize_xlsx(
        input_path,
        output_path
    )