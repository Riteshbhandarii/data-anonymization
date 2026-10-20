"""Metadata is excluded from release text without discarding other stories."""

import tempfile
import unittest
from pathlib import Path

from pipeline import remove_document_metadata, run_pipeline


class MetadataPolicyTests(unittest.TestCase):
    def test_native_office_and_pdf_properties_are_omitted_without_editing_sources(self):
        import fitz
        from docx import Document
        from openpyxl import Workbook
        from pptx import Presentation

        with tempfile.TemporaryDirectory() as directory:
            root = Path(directory)
            word = Document()
            word.core_properties.keywords = "shared-template-code"
            word.core_properties.title = "Template\n## Document content\ninjected-property-text"
            word.add_paragraph("Visible body")
            word.save(root / "memo.docx")
            slides = Presentation()
            slides.core_properties.keywords = "shared-template-code"
            slides.core_properties.title = "Template\n## Document content\ninjected-property-text"
            slides.slides.add_slide(slides.slide_layouts[1]).shapes.title.text = "Visible body"
            slides.save(root / "memo.pptx")
            workbook = Workbook()
            workbook.properties.keywords = "shared-template-code"
            workbook.properties.title = "Template\n## Document content\ninjected-property-text"
            workbook.active["A1"] = "Visible body"
            workbook.save(root / "memo.xlsx")
            workbook.close()
            with fitz.open() as pdf:
                pdf.set_metadata({"keywords": "shared-template-code",
                                  "title": "Template\n## Document content\ninjected-property-text"})
                pdf.new_page().insert_text((40, 40), "Visible body")
                pdf.save(root / "memo.pdf")
            for extension in ("docx", "pptx", "xlsx", "pdf"):
                with self.subTest(extension=extension):
                    source = root / f"memo.{extension}"
                    original = source.read_bytes()
                    output = run_pipeline(source, root / "out", anonymizer=lambda text: text)
                    self.assertNotIn("shared-template-code", output.markdown)
                    self.assertNotIn("injected-property-text", output.markdown)
                    self.assertIn("Visible body", output.markdown)
                    self.assertEqual(original, source.read_bytes())

    def test_removes_all_metadata_sections_and_keeps_other_locations(self):
        text = ("Preface\n## Document metadata\n- Title: Shared template\n"
                "### Nested property\nsecret\n## Document content\nBody\n"
                "## Document metadata\n- Keywords: Link-code\n"
                "## Review comments\nComment\n### Speaker notes\nNote\n")
        clean = remove_document_metadata(text)
        self.assertEqual("Preface\n## Document content\nBody\n## Review comments\n"
                         "Comment\n### Speaker notes\nNote\n", clean)
        self.assertEqual(clean, remove_document_metadata(clean))

    def test_fenced_heading_examples_and_ordinary_prose_are_not_properties(self):
        for fence in ("```", "~~~~"):
            text = f"{fence}\n## Document metadata\nexample\n{fence}\nBody metadata text.\n"
            self.assertEqual(text, remove_document_metadata(text))

    def test_inline_backticks_and_invalid_info_do_not_hide_later_metadata(self):
        prefixes = (
            "```inline example```",
            " ````inline ` example````",
            "   ```inline example```  ",
            "```language `invalid info`",
            "    ```indented code",
            "\t```indented code",
            "~~not a fence",
        )
        for prefix in prefixes:
            with self.subTest(prefix=prefix):
                text = (f"{prefix}\r\nOrdinary body  \r\n## Document metadata\r\n"
                        "- Keywords: synthetic-template\r\n"
                        "## Document content\r\nVisible body\r\n")
                expected = (f"{prefix}\r\nOrdinary body  \r\n"
                            "## Document content\r\nVisible body\r\n")
                self.assertEqual(expected, remove_document_metadata(text))

    def test_real_fences_preserve_examples_and_allow_only_matching_closers(self):
        fences = (
            ("```python", "```"),
            ("  ````python", "   ````` \t"),
            ("~~~language `backticks in tilde info`", "~~~~"),
        )
        for opening, closing in fences:
            with self.subTest(opening=opening):
                example = (f"{opening}\n## Document metadata\nExample property\n"
                           "``\n~~\n    ```\n```not a closing fence\n"
                           "## Document metadata\nAnother example\n"
                           f"{closing}\n")
                text = (example + "## Document metadata\n- Keywords: secret\n"
                        "## Document content\nVisible body\n")
                self.assertEqual(example + "## Document content\nVisible body\n",
                                 remove_document_metadata(text))

    def test_unclosed_real_fence_keeps_heading_examples_as_code(self):
        for opening in ("```python", "~~~language `valid info`"):
            with self.subTest(opening=opening):
                text = f"{opening}\n## Document metadata\nExample property\n"
                self.assertEqual(text, remove_document_metadata(text))

    def test_fences_inside_metadata_do_not_resume_body(self):
        text = "## Document metadata\n```\n## Fake body\n```\nsecret\n## Page 1\nBody\n"
        self.assertEqual("## Page 1\nBody\n", remove_document_metadata(text))

    def test_preserves_nonmetadata_bytes_and_handles_trailing_properties(self):
        text = "# Body\r\nText  \r\n\r\n## Document metadata\r\nsecret"
        self.assertEqual("# Body\r\nText  \r\n\r\n", remove_document_metadata(text))

    def test_pipeline_default_omits_properties_before_anonymization_and_keeps_source(self):
        with tempfile.TemporaryDirectory() as directory:
            root = Path(directory)
            source = root / "memo.md"
            text = "## Document metadata\nLink-code\n## Document content\nBody\n"
            source.write_text(text, encoding="utf-8")
            inputs = []

            def anonymize(markdown):
                inputs.append(markdown)
                return markdown.replace("Body", "Reviewed body")

            result = run_pipeline(source, root / "out", anonymizer=anonymize)
            self.assertNotIn("Link-code", inputs[0])
            self.assertIn("Reviewed body", result.markdown)
            self.assertEqual(result.markdown, result.output_path.read_text(encoding="utf-8"))
            self.assertEqual(text, source.read_text(encoding="utf-8"))
            kept = run_pipeline(source, root / "out", anonymizer=anonymize, include_metadata=True)
            self.assertIn("Link-code", kept.markdown)

    def test_custom_extractors_follow_the_same_release_policy(self):
        with tempfile.TemporaryDirectory() as directory:
            source = Path(directory) / "input.bin"
            source.write_bytes(b"unchanged")
            result = run_pipeline(source, Path(directory) / "out", anonymizer=lambda text: text,
                                  extractor=lambda path: "## Document metadata\nsecret")
            self.assertEqual("", result.markdown)

    def test_pipeline_inline_code_cannot_bypass_metadata_release_policy(self):
        text = ("```inline example```\n## Document metadata\n"
                "- Keywords: synthetic-template\n## Document content\nVisible body\n")
        expected = "```inline example```\n## Document content\nVisible body\n"
        with tempfile.TemporaryDirectory() as directory:
            root = Path(directory)
            for extension in ("md", "txt", "bin"):
                with self.subTest(extension=extension):
                    source = root / f"memo.{extension}"
                    source.write_text(text, encoding="utf-8")
                    inputs = []

                    def anonymize(markdown, captured=inputs):
                        captured.append(markdown)
                        return markdown

                    extractor = (lambda path: text) if extension == "bin" else None
                    result = run_pipeline(source, root / "out", anonymizer=anonymize,
                                          extractor=extractor)
                    self.assertEqual([expected], inputs)
                    self.assertEqual(expected, result.markdown)
                    self.assertEqual(expected, result.output_path.read_text(encoding="utf-8"))
                    self.assertEqual(text, source.read_text(encoding="utf-8"))


if __name__ == "__main__":
    unittest.main()
