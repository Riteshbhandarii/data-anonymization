"""Regression coverage for metadata, document stories, image OCR, and formulas."""

import shutil
import unittest
import xml.etree.ElementTree as ET
from io import BytesIO
from pathlib import Path
from tempfile import TemporaryDirectory
from unittest.mock import patch
from zipfile import ZipFile

from pipeline import ExtractionError, extract_to_markdown
from pipeline.validation import _text_by_location


class ExtractionGapTests(unittest.TestCase):
    def setUp(self):
        self.directory = TemporaryDirectory()
        self.addCleanup(self.directory.cleanup)
        self.root = Path(self.directory.name)

    def test_nested_docx_table_pipes_are_escaped_once(self):
        from docx import Document

        document = Document()
        outer = document.add_table(rows=1, cols=1)
        inner = outer.cell(0, 0).add_table(rows=1, cols=1)
        inner.cell(0, 0).text = "Aino | Example"
        path = self.root / "nested-pipes.docx"
        document.save(path)
        text = _text_by_location(extract_to_markdown(path), path)["body"]
        self.assertIn("Aino | Example", text)

    def test_docx_metadata_and_all_header_footer_variants(self):
        from docx import Document
        from docx.enum.section import WD_SECTION_START
        from docx.shared import Inches

        document = Document()
        document.add_paragraph("Body Contact")
        document.core_properties.author = "Metadata Author"
        document.core_properties.last_modified_by = "Metadata Editor"
        document.core_properties.keywords = "metadata-only@example.test"
        document.core_properties.comments = "Metadata Comment"
        first = document.sections[0]
        for name in (
            "header", "first_page_header", "even_page_header",
            "footer", "first_page_footer", "even_page_footer",
        ):
            story = getattr(first, name)
            story.paragraphs[0].text = f"Contact in {name}"
            story.add_table(rows=1, cols=1, width=Inches(4)).cell(0, 0).text = (
                f"Table contact in {name}"
            )
        # The second section inherits the same stories and must not repeat them.
        document.add_section(WD_SECTION_START.NEW_PAGE)
        path = self.root / "stories.docx"
        document.save(path)

        markdown = extract_to_markdown(path)
        self.assertIn("## Document metadata", markdown)
        self.assertIn("## Document content", markdown)
        for value in (
            "Body Contact", "Metadata Author", "Metadata Editor",
            "metadata-only@example.test", "Metadata Comment",
        ):
            self.assertIn(value, markdown)
        for name in (
            "header", "first_page_header", "even_page_header",
            "footer", "first_page_footer", "even_page_footer",
        ):
            self.assertEqual(markdown.count(f"Contact in {name}"), 1)
            self.assertEqual(markdown.count(f"Table contact in {name}"), 1)

    def test_docx_nested_tables_and_content_controls_keep_their_text(self):
        from docx import Document
        from docx.oxml import parse_xml

        document = Document()
        outer = document.add_table(rows=1, cols=1).cell(0, 0)
        outer.text = "Outer Contact"
        outer.add_table(rows=1, cols=1).cell(0, 0).text = "nested-only@example.test"
        document.element.body.insert(0, parse_xml(
            '<w:sdt xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main">'
            '<w:sdtContent><w:p><w:r><w:t>controlled-only@example.test</w:t></w:r></w:p>'
            '</w:sdtContent></w:sdt>'
        ))
        path = self.root / "nested.docx"
        document.save(path)
        markdown = extract_to_markdown(path)
        for value in ("Outer Contact", "nested-only@example.test", "controlled-only@example.test"):
            self.assertIn(value, markdown)

    def test_docx_body_table_and_header_images_reach_ocr(self):
        from docx import Document
        from docx.shared import Inches
        from PIL import Image

        image = self.root / "picture.png"
        Image.new("RGB", (200, 80), "white").save(image)
        document = Document()
        document.add_picture(str(image), width=Inches(2))
        document.add_table(rows=1, cols=1).cell(0, 0).paragraphs[0].add_run().add_picture(str(image))
        document.sections[0].header.paragraphs[0].add_run().add_picture(str(image))
        path = self.root / "pictures.docx"
        document.save(path)
        with patch("pipeline.extractors._ocr_image_bytes", side_effect=[
            "body-image@example.test", "table-image@example.test", "header-image@example.test",
        ]) as ocr:
            sections = _text_by_location(extract_to_markdown(path, "eng+fin"), path)
        self.assertEqual(3, ocr.call_count)
        self.assertTrue(all(call.args[1] == "eng+fin" for call in ocr.call_args_list))
        for value in ("body-image@example.test", "table-image@example.test", "header-image@example.test"):
            self.assertIn(value, sections["embedded_image"])

    def _docx_with_note_parts(self, *, broken_reference=False):
        from docx import Document
        from docx.opc.constants import RELATIONSHIP_TYPE as RT
        from docx.opc.packuri import PackURI
        from docx.opc.part import Part
        from docx.oxml import parse_xml

        document = Document()
        paragraph = document.add_paragraph("Body Contact")
        for kind, reltype in (("footnote", RT.FOOTNOTES), ("endnote", RT.ENDNOTES)):
            for identifier in ("2", "1", "2"):
                paragraph._p.append(parse_xml(
                    '<w:r xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main">'
                    f'<w:{kind}Reference w:id="{identifier}"/></w:r>'
                ))
            note_ids = ("1", "orphan") if broken_reference else ("1", "2", "orphan")
            content = ''.join(
                f'<w:{kind} w:id="{identifier}"><w:p><w:r><w:t>'
                f'{kind}-{identifier}@example.test</w:t></w:r></w:p></w:{kind}>'
                for identifier in note_ids
            )
            part = Part(
                PackURI(f"/word/{kind}s.xml"),
                f"application/vnd.openxmlformats-officedocument.wordprocessingml.{kind}s+xml",
                ('<w:' + kind + 's xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main">'
                 + content + f'</w:{kind}s>').encode(), document.part.package,
            )
            document.part.relate_to(part, reltype)
        path = self.root / "notes.docx"
        document.save(path)
        return path

    def test_docx_footnotes_and_endnotes_follow_references_and_have_notes_location(self):
        path = self._docx_with_note_parts()
        markdown = extract_to_markdown(path)
        sections = _text_by_location(markdown, path)
        self.assertIn("Body Contact", sections["body"])
        for kind in ("footnote", "endnote"):
            self.assertIn(f"[{kind}:2][{kind}:1][{kind}:2]", sections["body"])
            for identifier in ("1", "2"):
                value = f"{kind}-{identifier}@example.test"
                self.assertEqual(1, markdown.count(value))
                self.assertIn(value, sections["notes"])
                self.assertNotIn(value, sections["body"])
            self.assertLess(markdown.index(f"{kind}-2@"), markdown.index(f"{kind}-1@"))
            self.assertNotIn(f"{kind}-orphan@", markdown)

    def test_docx_missing_referenced_note_fails_instead_of_losing_text(self):
        path = self._docx_with_note_parts(broken_reference=True)
        with self.assertRaisesRegex(ExtractionError, "missing footnote 2"):
            extract_to_markdown(path)

    def test_docx_missing_note_part_fails_with_reference_context(self):
        from docx import Document
        from docx.oxml import parse_xml

        document = Document()
        document.add_paragraph("Body Contact")._p.append(parse_xml(
            '<w:r xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main">'
            '<w:footnoteReference w:id="1"/></w:r>'
        ))
        path = self.root / "missing-notes.docx"
        document.save(path)
        with self.assertRaisesRegex(ExtractionError, "references footnotes.*no footnotes part"):
            extract_to_markdown(path)

    def test_multiframe_tiff_routes_every_page_to_ocr(self):
        from PIL import Image

        path = self.root / "pages.tiff"
        Image.new("RGB", (200, 80), "white").save(
            path, save_all=True, append_images=[Image.new("RGB", (200, 80), "black")]
        )
        with patch("pipeline.extractors._run_tesseract", side_effect=[
            "first-page@example.test", "second-page@example.test",
        ]) as ocr:
            markdown = extract_to_markdown(path, "eng+fin")
        self.assertEqual(2, ocr.call_count)
        self.assertIn("first-page@example.test", markdown)
        self.assertIn("second-page@example.test", markdown)

    def test_pptx_metadata_and_speaker_notes(self):
        from pptx import Presentation

        presentation = Presentation()
        slide = presentation.slides.add_slide(presentation.slide_layouts[1])
        slide.shapes.title.text = "Body slide"
        slide.notes_slide.notes_text_frame.text = "Notes Contact"
        presentation.core_properties.author = "Slides Author"
        presentation.core_properties.last_modified_by = "Slides Editor"
        presentation.core_properties.keywords = "slides-only@example.test"
        path = self.root / "metadata.pptx"
        presentation.save(path)

        markdown = extract_to_markdown(path)
        for value in (
            "Slides Author", "Slides Editor", "slides-only@example.test",
            "### Speaker notes", "Notes Contact",
        ):
            self.assertIn(value, markdown)

    def test_docx_review_comments_and_revisions_have_distinct_locations(self):
        from docx import Document
        from docx.opc.constants import RELATIONSHIP_TYPE as RT
        from docx.opc.packuri import PackURI
        from docx.opc.part import Part
        from docx.oxml import parse_xml

        document = Document()
        paragraph = document.add_paragraph("Body Contact ")
        paragraph._p.append(parse_xml(
            '<w:ins xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main" '
            'w:id="1" w:author="Revision Editor"><w:r><w:t>Inserted Person</w:t></w:r></w:ins>'
        ))
        paragraph._p.append(parse_xml(
            '<w:del xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main" '
            'w:id="2" w:author="Deletion Editor"><w:r><w:delText>Deleted Person</w:delText>'
            '</w:r></w:del>'
        ))
        comments = Part(
            PackURI("/word/comments.xml"),
            "application/vnd.openxmlformats-officedocument.wordprocessingml.comments+xml",
            b'<w:comments xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main">'
            b'<w:comment w:id="0" w:author="Comment Author" w:initials="CA">'
            b'<w:p><w:r><w:t>review-only@example.test</w:t></w:r></w:p>'
            b'</w:comment></w:comments>',
            document.part.package,
        )
        document.part.relate_to(comments, RT.COMMENTS)
        path = self.root / "review.docx"
        document.save(path)
        sections = _text_by_location(extract_to_markdown(path), path)
        self.assertIn("Body Contact Inserted Person", sections["body"])
        self.assertNotIn("Deleted Person", sections["body"])
        for value in ("Inserted Person", "Revision Editor", "Deleted Person", "Deletion Editor"):
            self.assertIn(value, sections["tracked_change"])
        for value in ("Comment Author", "review-only@example.test"):
            self.assertIn(value, sections["comment"])
            self.assertNotIn(value, sections["body"])

    def test_pptx_nested_picture_ocr_keeps_body_and_notes_separate(self):
        from PIL import Image
        from pptx import Presentation
        from pptx.util import Inches

        image_path = self.root / "image.png"
        Image.new("RGB", (640, 160), "white").save(image_path)
        presentation = Presentation()
        slide = presentation.slides.add_slide(presentation.slide_layouts[5])
        slide.shapes.title.text = "Body Contact"
        group = slide.shapes.add_group_shape()
        group.shapes.add_picture(str(image_path), 0, 0, width=Inches(4))
        slide.notes_slide.notes_text_frame.text = "Notes Contact"
        path = self.root / "group.pptx"
        presentation.save(path)
        with patch("pipeline.extractors._ocr_image_bytes", return_value="Image Contact") as ocr:
            sections = _text_by_location(extract_to_markdown(path, "eng+fin"), path)
        ocr.assert_called_once()
        self.assertEqual("eng+fin", ocr.call_args.args[1])
        self.assertIn("Image Contact", sections["embedded_image"])
        self.assertNotIn("Image Contact", sections["body"])
        self.assertIn("Body Contact", sections["body"])
        self.assertIn("Notes Contact", sections["notes"])

    def test_pptx_unembedded_picture_keeps_text_notes_and_other_images(self):
        from PIL import Image
        from pptx import Presentation
        from pptx.opc.constants import RELATIONSHIP_TYPE as RT
        from pptx.oxml.ns import qn
        from pptx.util import Inches

        pixels = BytesIO()
        Image.new("RGB", (200, 80), "white").save(pixels, format="PNG")
        for linked in (False, True):
            with self.subTest(linked=linked):
                presentation = Presentation()
                slide = presentation.slides.add_slide(presentation.slide_layouts[5])
                slide.shapes.title.text = "Body Contact"
                slide.shapes.add_textbox(0, 0, Inches(4), Inches(1)).text = "Slide content"
                picture = slide.shapes.add_picture(pixels, 0, Inches(1))
                blip = picture._pic.blipFill.blip
                del blip.attrib[qn("r:embed")]
                if linked:
                    relation = slide.part.relate_to(
                        "https://example.invalid/image.png", RT.IMAGE, is_external=True
                    )
                    blip.set(qn("r:link"), relation)
                slide.shapes.add_picture(pixels, 0, Inches(2))
                slide.notes_slide.notes_text_frame.text = "Notes Contact"
                path = self.root / "unembedded.pptx"
                presentation.save(path)

                with patch("pipeline.extractors._ocr_image_bytes", return_value="Image Contact") as ocr:
                    markdown = extract_to_markdown(path, "eng+fin")
                ocr.assert_called_once()
                self.assertEqual("eng+fin", ocr.call_args.args[1])
                for value in ("Body Contact", "Slide content", "Notes Contact", "Image Contact"):
                    self.assertIn(value, markdown)
                self.assertIn("Image not extracted", markdown)

    def test_pptx_picture_placeholder_still_uses_ocr(self):
        from PIL import Image
        from pptx import Presentation

        pixels = BytesIO()
        Image.new("RGB", (200, 80), "white").save(pixels, format="PNG")
        presentation = Presentation()
        slide = presentation.slides.add_slide(presentation.slide_layouts[8])
        slide.placeholders[1].insert_picture(pixels)
        path = self.root / "picture-placeholder.pptx"
        presentation.save(path)

        with patch("pipeline.extractors._ocr_image_bytes", return_value="Placeholder Contact") as ocr:
            markdown = extract_to_markdown(path)
        ocr.assert_called_once()
        self.assertIn("Placeholder Contact", markdown)
        self.assertNotIn("Image not extracted", markdown)

    def test_standalone_image_routes_pixels_to_ocr(self):
        from PIL import Image

        path = self.root / "standalone.png"
        Image.new("RGB", (320, 100), "white").save(path)
        with patch("pipeline.extractors._run_tesseract", return_value="Image Contact") as ocr:
            markdown = extract_to_markdown(path, "eng+fin")
        ocr.assert_called_once()
        self.assertEqual("eng+fin", ocr.call_args.args[1])
        self.assertIn("Image Contact", _text_by_location(markdown, path)["body"])

    @unittest.skipUnless(shutil.which("tesseract"), "Tesseract is not installed")
    def test_standalone_png_real_ocr_reads_identifier_from_pixels(self):
        import pymupdf

        # The source file contains pixels only; no label sidecar or mock can
        # supply the expected identifier to the standalone image reader.
        identifier = "standalone-image@example.test"
        path = self.root / "standalone-text.png"
        with pymupdf.open() as raster_source:
            page = raster_source.new_page(width=480, height=100)
            page.insert_text((20, 60), identifier, fontsize=24)
            page.get_pixmap(matrix=pymupdf.Matrix(2, 2)).save(path)

        markdown = extract_to_markdown(path, ocr_language="eng")
        self.assertIn(identifier, _text_by_location(markdown, path)["body"])

    def test_xlsx_metadata_and_hidden_worksheet(self):
        from openpyxl import Workbook

        workbook = Workbook()
        workbook.active.append(["Visible Contact"])
        hidden = workbook.create_sheet("Private")
        hidden.sheet_state = "hidden"
        hidden.append(["Hidden Contact"])
        workbook.properties.creator = "Workbook Author"
        workbook.properties.lastModifiedBy = "Workbook Editor"
        workbook.properties.keywords = "workbook-only@example.test"
        path = self.root / "metadata.xlsx"
        workbook.save(path)
        workbook.close()

        markdown = extract_to_markdown(path)
        for value in (
            "Workbook Author", "Workbook Editor", "workbook-only@example.test",
            "## Sheet: Private", "Hidden Contact", "Visible Contact",
        ):
            self.assertIn(value, markdown)

    def test_xlsx_comments_on_blank_cells_and_hidden_sheets_are_returned(self):
        from openpyxl import Workbook
        from openpyxl.comments import Comment

        workbook = Workbook()
        workbook.active["A1"] = "Body Contact"
        workbook.active["C3"].comment = Comment("blank-cell@example.test", "Visible Author")
        hidden = workbook.create_sheet("Private")
        hidden.sheet_state = "hidden"
        hidden["A1"].comment = Comment("hidden-comment@example.test", "Hidden Author")
        path = self.root / "comments.xlsx"
        workbook.save(path)
        workbook.close()
        sections = _text_by_location(extract_to_markdown(path), path)
        for value in ("blank-cell@example.test", "Visible Author", "hidden-comment@example.test", "Hidden Author"):
            self.assertIn(value, sections["comment"])
            self.assertNotIn(value, sections["body"])

    def test_xlsx_incorrect_stored_dimensions_do_not_truncate_values(self):
        from openpyxl import Workbook

        workbook = Workbook()
        workbook.active["A1"] = "First Contact"
        workbook.active["C3"] = "last-contact@example.test"
        path = self.root / "dimensions.xlsx"
        workbook.save(path)
        workbook.close()
        with ZipFile(path) as archive:
            entries = [(item, archive.read(item.filename)) for item in archive.infolist()]
        with ZipFile(path, "w") as archive:
            for item, data in entries:
                if item.filename == "xl/worksheets/sheet1.xml":
                    data = data.replace(b'ref="A1:C3"', b'ref="A1:A1"')
                archive.writestr(item, data)
        self.assertIn("last-contact@example.test", extract_to_markdown(path))

    def test_pdf_metadata_and_text_without_unnecessary_ocr(self):
        import pymupdf

        path = self.root / "text.pdf"
        with pymupdf.open() as document:
            document.new_page().insert_text((40, 40), "Visible Contact")
            document.set_metadata({
                "author": "PDF Author",
                "subject": "pdf-only@example.test",
                "keywords": "PDF Keywords",
            })
            document.save(path)

        with patch("pipeline.extractors._ocr_image_bytes") as ocr:
            markdown = extract_to_markdown(path)
        ocr.assert_not_called()
        for value in ("PDF Author", "pdf-only@example.test", "PDF Keywords", "Visible Contact"):
            self.assertIn(value, markdown)

    def test_pdf_annotation_content_and_author_have_comment_location(self):
        import pymupdf

        path = self.root / "annotations.pdf"
        with pymupdf.open() as document:
            page = document.new_page()
            page.insert_text((40, 40), "Visible Contact")
            annotation = page.add_text_annot((80, 80), "annotation-only@example.test")
            annotation.set_info(title="Private Reviewer", subject="Private Subject")
            annotation.update()
            document.save(path)
        with patch("pipeline.extractors._ocr_image_bytes") as ocr:
            sections = _text_by_location(extract_to_markdown(path), path)
        ocr.assert_not_called()
        for value in ("annotation-only@example.test", "Private Reviewer", "Private Subject"):
            self.assertIn(value, sections["comment"])
            self.assertNotIn(value, sections["body"])

    def _mixed_pdf(self, rotation: int = 0) -> Path:
        import pymupdf

        # Rasterize known text before embedding it; the final PDF has no text
        # objects for this identifier, so it can only be returned through OCR.
        with pymupdf.open() as image_document:
            image_page = image_document.new_page(width=480, height=100)
            image_page.insert_text((20, 60), "image-only@example.test", fontsize=24)
            image_bytes = image_page.get_pixmap(matrix=pymupdf.Matrix(2, 2)).tobytes("png")
        path = self.root / "mixed.pdf"
        with pymupdf.open() as document:
            page = document.new_page(width=600, height=400)
            page.insert_text((40, 40), "Visible Contact")
            page.insert_image(pymupdf.Rect(40, 100, 520, 200), stream=image_bytes)
            page.set_rotation(rotation)
            document.save(path)
        return path

    def test_mixed_pdf_ocr_keeps_image_evidence_separate_from_text_layer(self):
        from PIL import Image

        for rotation, size in ((0, (960, 200)), (90, (200, 960))):
            with self.subTest(rotation=rotation):
                path = self._mixed_pdf(rotation)
                with patch(
                    "pipeline.extractors._ocr_image_bytes",
                    return_value="Visible Contact\nimage-only@example.test\nimage-only@example.test",
                ) as ocr:
                    markdown = extract_to_markdown(path, ocr_language="eng")
                ocr.assert_called_once()
                self.assertEqual(ocr.call_args.args[1], "eng")
                with Image.open(BytesIO(ocr.call_args.args[0])) as region:
                    # Only the image region is passed to OCR, not the entire page.
                    self.assertEqual(region.size, size)
                sections = _text_by_location(markdown, path)
                self.assertIn("Visible Contact", sections["body"])
                self.assertNotIn("image-only@example.test", sections["body"])
                self.assertIn("Visible Contact", sections["embedded_image"])
                self.assertEqual(markdown.count("image-only@example.test"), 1)

    def test_scanned_pdf_still_uses_full_page_ocr(self):
        import pymupdf

        path = self.root / "scanned.pdf"
        with pymupdf.open() as document:
            document.new_page()
            document.save(path)
        with patch(
            "pipeline.extractors._ocr_image_bytes", return_value="Scanned Contact"
        ) as ocr:
            markdown = extract_to_markdown(path)
        ocr.assert_called_once()
        self.assertIn("Scanned Contact", markdown)

    @unittest.skipUnless(shutil.which("tesseract"), "Tesseract is not installed")
    def test_mixed_pdf_real_ocr_reads_identifier_inside_image(self):
        markdown = extract_to_markdown(self._mixed_pdf())
        self.assertIn("Visible Contact", markdown)
        self.assertIn("image-only@example.test", markdown)

    @unittest.skipUnless(shutil.which("tesseract"), "Tesseract is not installed")
    def test_pptx_real_ocr_reads_identifier_only_present_in_pixels(self):
        import pymupdf
        from pptx import Presentation
        from pptx.util import Inches

        # This is an independent known-text raster, not a generator label reader.
        with pymupdf.open() as image_document:
            page = image_document.new_page(width=480, height=100)
            page.insert_text((20, 60), "slide-image@example.test", fontsize=24)
            image_bytes = page.get_pixmap(matrix=pymupdf.Matrix(2, 2)).tobytes("png")
        presentation = Presentation()
        slide = presentation.slides.add_slide(presentation.slide_layouts[5])
        slide.shapes.title.text = "Visible slide"
        slide.shapes.add_picture(BytesIO(image_bytes), Inches(1), Inches(2), width=Inches(5))
        path = self.root / "image-text.pptx"
        presentation.save(path)
        sections = _text_by_location(extract_to_markdown(path), path)
        self.assertIn("slide-image@example.test", sections["embedded_image"])
        self.assertNotIn("slide-image@example.test", sections["body"])

    def _formula_workbook(self) -> Path:
        from openpyxl import Workbook

        workbook = Workbook()
        workbook.active.title = "Contacts"
        workbook.active.append(["Computed name"])
        workbook.active["A2"] = '=CONCAT("Formula", " Contact")'
        path = self.root / "formulas.xlsx"
        workbook.save(path)
        workbook.close()
        return path

    def _set_formula_cache(self, path: Path, value: str, result_type: str = "str"):
        # openpyxl writes formulas but does not calculate their cached values.
        # Patch the OOXML cache to model a workbook saved by Excel/LibreOffice.
        with ZipFile(path) as archive:
            entries = [(item, archive.read(item.filename)) for item in archive.infolist()]
        namespace = {"s": "http://schemas.openxmlformats.org/spreadsheetml/2006/main"}
        with ZipFile(path, "w") as archive:
            for item, data in entries:
                if item.filename == "xl/worksheets/sheet1.xml":
                    root = ET.fromstring(data)
                    cell = root.find(".//s:c[@r='A2']", namespace)
                    cell.set("t", result_type)
                    cached = cell.find("s:v", namespace)
                    cached.text = value
                    data = ET.tostring(root, encoding="utf-8", xml_declaration=True)
                archive.writestr(item, data)

    def test_xlsx_formula_keeps_computed_result_and_expression(self):
        path = self._formula_workbook()
        self._set_formula_cache(path, "Formula Contact")
        markdown = extract_to_markdown(path)
        self.assertIn("Formula Contact", markdown)
        self.assertIn('=CONCAT("Formula", " Contact")', markdown)

    def test_xlsx_formula_literal_is_retained_when_absent_from_computed_result(self):
        from openpyxl import Workbook

        workbook = Workbook()
        workbook.active.title = "Contacts"
        workbook.active["A1"] = "Result"
        workbook.active["A2"] = '=IF(FALSE,"formula-only@example.test","Public result")'
        path = self.root / "private-formula.xlsx"
        workbook.save(path)
        workbook.close()
        self._set_formula_cache(path, "Public result")
        sections = _text_by_location(extract_to_markdown(path), path)
        self.assertIn("Public result", sections["body"])
        self.assertIn("formula-only@example.test", sections["body"])

    def test_xlsx_formula_without_cache_fails_with_cell_context(self):
        path = self._formula_workbook()
        with self.assertRaisesRegex(ExtractionError, "Contacts.*A2.*cached result"):
            extract_to_markdown(path)

    def test_xlsx_formula_with_cached_empty_string_is_valid(self):
        path = self._formula_workbook()
        self._set_formula_cache(path, "")
        markdown = extract_to_markdown(path)
        self.assertIn("Computed name", markdown)
        self.assertIn('=CONCAT("Formula", " Contact")', markdown)

    def test_xlsx_formula_error_cache_is_not_treated_as_content(self):
        path = self._formula_workbook()
        self._set_formula_cache(path, "#REF!", "e")
        with self.assertRaisesRegex(ExtractionError, "Recalculate and save"):
            extract_to_markdown(path)


if __name__ == "__main__":
    unittest.main()
