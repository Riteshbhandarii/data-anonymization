"""Keep full-document detection scoring distinct from body-only evaluation."""

import csv
import json
import tempfile
import unittest
from pathlib import Path

from docx import Document

from eval.bench import fingerprint
from eval.pipeline_bench import (
    evaluate_pipeline,
    occurrences,
    score_label,
    write_reports,
)
from redact import Span


class ExactSyntheticDetector:
    """Test fixture detector using only known fictional names, not production logic."""

    def __init__(self):
        self.inputs = []

    def detect(self, text, language="en"):
        self.inputs.append((text, language))
        result = []
        for value in ("Alice Example", "Bob Example", "Carol Example"):
            start = text.find(value)
            while start >= 0:
                result.append(Span(start, start + len(value), "PERSON", 1))
                start = text.find(value, start + 1)
        return result


class PipelineBenchmarkTests(unittest.TestCase):
    def make_corpus(self, directory):
        root = Path(directory)
        for name in ("labels", "text", "docs"):
            (root / name).mkdir()
        document = Document()
        document.add_paragraph("Alice Example wrote this.")
        document.sections[0].header.paragraphs[0].text = "Bob Example"
        document.core_properties.author = "Carol Example"
        document.save(root / "docs" / "sample.docx")
        (root / "text" / "sample.txt").write_text("Alice Example wrote this.", encoding="utf-8")
        labels = {
            "file": "docs/sample.docx", "language": "en", "format": "docx",
            "entities": [{"type": "PERSON", "value": name, "location": location}
                         for name, location in (("Alice Example", "body"), ("Bob Example", "header"),
                                                ("Carol Example", "metadata"))],
        }
        (root / "labels" / "sample.json").write_text(json.dumps(labels), encoding="utf-8")
        with (root / "index.csv").open("w", encoding="utf-8", newline="") as handle:
            writer = csv.DictWriter(handle, fieldnames=["file", "language", "format"])
            writer.writeheader()
            writer.writerow({key: labels[key] for key in writer.fieldnames})
        (root / "corpus.json").write_text(json.dumps({"corpus_sha256": fingerprint(root, ["sample"])}), encoding="utf-8")
        return root

    def test_actual_document_path_includes_header_and_metadata_not_in_sidecar(self):
        with tempfile.TemporaryDirectory() as directory:
            root = self.make_corpus(directory)
            detector = ExactSyntheticDetector()
            report = evaluate_pipeline(root, detector, output_dir=root / "outputs")
            self.assertEqual(len(detector.inputs), 2)
            self.assertNotIn("Bob Example", detector.inputs[0][0])
            self.assertIn("Bob Example", detector.inputs[1][0])
            self.assertIn("Carol Example", detector.inputs[1][0])
            self.assertEqual([row["labels"] for row in report["body_comparison"]], [1, 1])
            self.assertEqual([row["whole"] for row in report["body_comparison"]], [1, 1])
            document = report["documents"][0]
            self.assertEqual(document["errors"], [])
            self.assertEqual(len(document["scores"]["full_pipeline"]), 3)
            self.assertTrue(all(entity["whole"] for entity in document["scores"]["full_pipeline"]))
            output = Path(document["output_path"]).read_text(encoding="utf-8")
            for value in ("Alice Example", "Bob Example", "Carol Example"):
                self.assertNotIn(value, output)
            write_reports(report, root / "report")
            self.assertTrue((root / "report" / "summary.md").is_file())
            self.assertTrue((root / "report" / "by_location.csv").is_file())

    def test_extraction_errors_keep_all_labels_in_full_path_denominator(self):
        def broken_extractor(_path):
            raise ValueError("fixture extraction failure")

        with tempfile.TemporaryDirectory() as directory:
            root = self.make_corpus(directory)
            report = evaluate_pipeline(root, ExactSyntheticDetector(), extractor=broken_extractor, output_dir=root / "out")
            full = report["documents"][0]["scores"]["full_pipeline"]
            self.assertEqual(len(full), 3)
            self.assertTrue(all(not entity["extracted"] and not entity["whole"] for entity in full))
            self.assertIn("fixture extraction failure", report["documents"][0]["errors"][0])
            self.assertIsNone(report["documents"][0]["output_path"])

    def test_missing_document_keeps_labels_and_records_missing_artifact(self):
        with tempfile.TemporaryDirectory() as directory:
            root = self.make_corpus(directory)
            (root / "docs" / "sample.docx").unlink()
            report = evaluate_pipeline(root, ExactSyntheticDetector(), output_dir=root / "out")
            document = report["documents"][0]
            self.assertIsNone(document["source_sha256"])
            self.assertEqual("missing", document["source_status"])
            self.assertIsNone(document["output_path"])
            self.assertTrue(document["errors"])
            self.assertEqual(3, len(document["scores"]["full_pipeline"]))
            self.assertTrue(all(not entity["extracted"] for entity in document["scores"]["full_pipeline"]))
            self.assertTrue(document["scores"]["detector_only"][0]["whole"])
            repeated = evaluate_pipeline(root, ExactSyntheticDetector(), output_dir=root / "out")
            self.assertEqual(report["source_artifacts_sha256"], repeated["source_artifacts_sha256"])

    def test_occurrence_offsets_survive_nfc_whitespace_and_markdown_escaping(self):
        text = "Prelude\n## Document metadata\nA\u030asa\\|Example<br>Person\n"
        value = "Åsa|Example Person"
        start = text.index("A\u030a")
        expected = (start, len(text) - 1)
        regions = [("metadata", text.index("##"), len(text))]
        self.assertEqual(occurrences(text, value, regions, "metadata"), [expected])
        label = {"type": "PERSON", "value": value, "location": "metadata"}
        scored = score_label(text, label, [Span(*expected, "PERSON", 1)], regions)
        self.assertTrue(scored["whole"])

    def test_matching_body_copy_cannot_satisfy_metadata(self):
        text = "Alice Example\n## Document metadata\nSomeone else\n"
        boundary = text.index("##")
        regions = [("body", 0, boundary), ("metadata", boundary, len(text))]
        label = {"type": "PERSON", "value": "Alice Example", "location": "metadata"}
        result = score_label(text, label, [Span(0, 13, "PERSON", 1)], regions)
        self.assertFalse(result["extracted"])
        self.assertFalse(result["whole"])

    def test_every_occurrence_and_whole_value_are_required(self):
        text = "Alice Example met Alice Example"
        entity = {"type": "PERSON", "value": "Alice Example", "location": "body"}
        ranges = [("body", 0, len(text))]
        result = score_label(text, entity, [Span(0, 13, "PERSON", 1), Span(18, 23, "PERSON", 1)], ranges, expected=2)
        self.assertFalse(result["whole"])
        self.assertTrue(result["partial"])
        self.assertEqual([item["detection"] for item in result["occurrences"]], ["whole", "partial"])
        # Losing a repeated occurrence during extraction cannot become 100% recall.
        result = score_label("Alice Example", entity, [Span(0, 13, "PERSON", 1)], [("body", 0, 13)], expected=2)
        self.assertFalse(result["whole"])
        self.assertTrue(result["extraction_incomplete"])
        self.assertEqual(result["missing_occurrences_minimum"], 1)

    def test_wrong_type_and_split_partial_spans_do_not_count_as_whole(self):
        entity = {"type": "PERSON", "value": "Alice Example", "location": "body"}
        ranges = [("body", 0, 13)]
        wrong = score_label("Alice Example", entity, [Span(0, 13, "LOCATION", 1)], ranges)
        self.assertFalse(wrong["whole"])
        self.assertFalse(wrong["partial"])
        split = score_label("Alice Example", entity, [Span(0, 5, "PERSON", 1), Span(6, 13, "PERSON", 1)], ranges)
        self.assertFalse(split["whole"])
        self.assertTrue(split["partial"])

    def test_fingerprint_mismatch_and_empty_manifest_fail_explicitly(self):
        with tempfile.TemporaryDirectory() as directory:
            root = self.make_corpus(directory)
            (root / "text" / "sample.txt").write_text("changed", encoding="utf-8")
            with self.assertRaisesRegex(ValueError, "fingerprint"):
                evaluate_pipeline(root, ExactSyntheticDetector(), output_dir=root / "out")
            (root / "index.csv").write_text("file,language,format\n", encoding="utf-8")
            with self.assertRaisesRegex(ValueError, "distinct documents"):
                evaluate_pipeline(root, ExactSyntheticDetector(), output_dir=root / "out")


if __name__ == "__main__":
    unittest.main()
