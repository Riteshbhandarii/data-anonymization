from pathlib import Path
from openpyxl import load_workbook


SENSITIVE_FIELDS = {
    "name": "PERSON",
    "email": "EMAIL",
    "phone": "PHONE",
    "company": "COMPANY",
    "iban": "IBAN",
    "date": "DATE",
    "personal_id": "PERSONAL_ID",
    "plate": "PLATE",
    "address": "ADDRESS",
}


def detect_sensitive_data(file_path):
    file_path = Path(file_path)

    workbook = load_workbook(file_path, data_only=False)

    detected = []

    print("=" * 70)
    print(f"FILE: {file_path.name}")
    print("=" * 70)

    # -------------------------
    # 1. Check metadata
    # -------------------------

    print("\n[METADATA DETECTION]")

    metadata = {
        "creator": workbook.properties.creator,
        "lastModifiedBy": workbook.properties.lastModifiedBy,
    }

    for field, value in metadata.items():

        if value:
            detected.append({
                "source": "metadata",
                "location": field,
                "type": "PERSON",
                "value": value
            })

            print(
                f"Detected PERSON | "
                f"{field} | {value}"
            )

    # -------------------------
    # 2. Check worksheets
    # -------------------------

    print("\n[CELL DETECTION]")

    for sheet in workbook.worksheets:

        headers = {}

        # Read first row as headers
        for cell in sheet[1]:

            if cell.value:

                header = str(cell.value).strip().lower()

                headers[cell.column] = header

        # Read remaining rows
        for row in sheet.iter_rows(min_row=2):

            for cell in row:

                if cell.value is None:
                    continue

                header = headers.get(cell.column)

                if header in SENSITIVE_FIELDS:

                    data_type = SENSITIVE_FIELDS[header]

                    detected.append({
                        "source": "cell",
                        "sheet": sheet.title,
                        "sheet_state": sheet.sheet_state,
                        "cell": cell.coordinate,
                        "type": data_type,
                        "value": cell.value
                    })

                    print(
                        f"{data_type:<12} | "
                        f"Sheet: {sheet.title:<30} | "
                        f"State: {sheet.sheet_state:<7} | "
                        f"Cell: {cell.coordinate:<4} | "
                        f"Value: {cell.value}"
                    )

    # -------------------------
    # Summary
    # -------------------------

    print("\n[SUMMARY]")

    counts = {}

    for item in detected:

        data_type = item["type"]

        counts[data_type] = counts.get(data_type, 0) + 1

    for data_type, count in counts.items():

        print(f"{data_type}: {count}")

    print(f"\nTotal sensitive values detected: {len(detected)}")

    return detected


if __name__ == "__main__":

    file_path = input(
        "Enter XLSX file path: "
    ).strip().strip('"')

    detect_sensitive_data(file_path)