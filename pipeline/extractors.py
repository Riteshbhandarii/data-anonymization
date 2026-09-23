"""Extract text from supported files and represent it as Markdown."""

import csv
from io import BytesIO
from pathlib import Path
from typing import Any

from .errors import ExtractionError, UnsupportedFormatError

TEXT_EXTENSIONS = {".md", ".txt"}
IMAGE_EXTENSIONS = {".bmp", ".jpeg", ".jpg", ".png", ".tif", ".tiff"}
SUPPORTED_EXTENSIONS = tuple(
    sorted(TEXT_EXTENSIONS | IMAGE_EXTENSIONS | {".csv", ".docx", ".pdf", ".pptx", ".xlsx"})
)


def detect_format(file_path: str | Path) -> str:
    """Return the normalized extension for a supported input file."""
    suffix = Path(file_path).suffix.lower()
    if suffix not in SUPPORTED_EXTENSIONS:
        supported = ", ".join(SUPPORTED_EXTENSIONS)
        raise UnsupportedFormatError(
            f"Unsupported file type '{suffix or '(none)'}'. Supported types: {supported}."
        )
    return suffix


def extract_to_markdown(file_path: str | Path, ocr_language: str = "eng") -> str:
    """Extract one supported file into a Markdown string."""
    path = Path(file_path)
    if not path.is_file():
        raise ExtractionError(f"Input file does not exist: {path}")

    suffix = detect_format(path)
    try:
        if suffix in TEXT_EXTENSIONS:
            return _extract_text(path)
        if suffix == ".csv":
            return _extract_csv(path)
        if suffix == ".docx":
            return _extract_docx(path)
        if suffix == ".xlsx":
            return _extract_xlsx(path)
        if suffix == ".pptx":
            return _extract_pptx(path)
        if suffix == ".pdf":
            return _extract_pdf(path, ocr_language)
        if suffix in IMAGE_EXTENSIONS:
            return _extract_image(path, ocr_language)
    except ExtractionError:
        raise
    except Exception as exc:
        raise ExtractionError(f"Could not extract '{path.name}': {exc}") from exc

    raise UnsupportedFormatError(f"No extractor registered for '{suffix}'.")


def _extract_text(path: Path) -> str:
    return path.read_text(encoding="utf-8-sig")


def _extract_csv(path: Path) -> str:
    with path.open("r", encoding="utf-8-sig", newline="") as handle:
        rows = list(csv.reader(handle))
    return _rows_to_markdown(rows, empty_message="_Empty CSV file._")


def _extract_docx(path: Path) -> str:
    from docx import Document

    document = Document(path)
    parts = _office_metadata(document.core_properties)
    parts.append("## Document content")
    parts.extend(_docx_blocks(document.element.body, document))

    # Linked sections reuse an earlier story; only extract its defining part.
    seen_parts: set[str] = set()
    for index, section in enumerate(document.sections, start=1):
        for name in (
            "header", "first_page_header", "even_page_header",
            "footer", "first_page_footer", "even_page_footer",
        ):
            story = getattr(section, name)
            if story.is_linked_to_previous:
                continue
            part_name = str(story.part.partname)
            if part_name in seen_parts:
                continue
            seen_parts.add(part_name)
            blocks = _docx_blocks(story._element, story)
            if blocks:
                label = name.replace("_", " ").capitalize()
                parts.append(f"## {label} (section {index})")
                parts.extend(blocks)

    return "\n\n".join(parts)


def _docx_blocks(element: Any, parent: Any) -> list[str]:
    """Read paragraphs and tables in body, header, or footer order."""
    from docx.table import Table
    from docx.text.paragraph import Paragraph

    parts: list[str] = []
    for child in element.iterchildren():
        if child.tag.endswith("}p"):
            paragraph = Paragraph(child, parent)
            text = paragraph.text.strip()
            if not text:
                continue

            style_name = paragraph.style.name if paragraph.style else ""
            if style_name.startswith("Heading"):
                level_text = style_name.removeprefix("Heading").strip()
                level = int(level_text) if level_text.isdigit() else 2
                parts.append(f"{'#' * max(1, min(level, 6))} {text}")
            elif style_name.startswith("List"):
                parts.append(f"- {text}")
            else:
                parts.append(text)
        elif child.tag.endswith("}tbl"):
            table = Table(child, parent)
            rows = [[cell.text for cell in row.cells] for row in table.rows]
            parts.append(_rows_to_markdown(rows, empty_message="_Empty table._"))

    return parts


def _extract_xlsx(path: Path) -> str:
    from openpyxl import load_workbook

    workbook = load_workbook(path, read_only=True, data_only=False)
    try:
        cached_workbook = load_workbook(path, read_only=True, data_only=True)
        try:
            properties = workbook.properties
            parts = _metadata_section({
                "Author": properties.creator,
                "Last modified by": properties.lastModifiedBy,
                "Title": properties.title,
                "Subject": properties.subject,
                "Description": properties.description,
                "Keywords": properties.keywords,
                "Category": properties.category,
                "Identifier": properties.identifier,
                "Language": properties.language,
                "Content status": properties.contentStatus,
                "Version": properties.version,
            })
            for sheet in workbook.worksheets:
                parts.append(f"## Sheet: {sheet.title}")
                cached_sheet = cached_workbook[sheet.title]
                rows: list[list[str]] = []
                for source_row, cached_row in zip(sheet.iter_rows(), cached_sheet.iter_rows()):
                    values = []
                    for source_cell, cached_cell in zip(source_row, cached_row):
                        value = cached_cell.value
                        # An explicitly stored empty string is a valid formula
                        # result; an untyped missing numeric cache is not.
                        missing_result = value is None and cached_cell.data_type != "str"
                        if source_cell.data_type == "f" and (
                            missing_result or cached_cell.data_type == "e"
                        ):
                            raise ExtractionError(
                                f"Formula in '{path.name}', sheet '{sheet.title}', "
                                f"cell {source_cell.coordinate} has no usable cached result. "
                                "Recalculate and save the workbook in a spreadsheet "
                                "application before extraction."
                            )
                        values.append("" if value is None else str(value))
                    if any(value.strip() for value in values):
                        rows.append(values)
                parts.append(_rows_to_markdown(rows, empty_message="_Empty sheet._"))
        finally:
            cached_workbook.close()
    finally:
        workbook.close()
    return "\n\n".join(parts)


def _extract_pptx(path: Path) -> str:
    from pptx import Presentation

    presentation = Presentation(path)
    parts = _office_metadata(presentation.core_properties)

    for index, slide in enumerate(presentation.slides, start=1):
        title_shape = slide.shapes.title
        title = title_shape.text.strip() if title_shape and title_shape.has_text_frame else ""
        heading = f"## Slide {index}"
        if title:
            heading = f"{heading}: {title}"
        parts.append(heading)

        for shape in slide.shapes:
            if title_shape is not None and shape.shape_id == title_shape.shape_id:
                continue
            if getattr(shape, "has_table", False):
                rows = [[cell.text for cell in row.cells] for row in shape.table.rows]
                parts.append(_rows_to_markdown(rows, empty_message="_Empty table._"))
            elif getattr(shape, "has_text_frame", False):
                text = shape.text.strip()
                if text:
                    parts.append(text)

        if slide.has_notes_slide:
            notes = slide.notes_slide.notes_text_frame.text.strip()
            if notes:
                parts.append(f"### Speaker notes\n\n{notes}")

    return "\n\n".join(parts)


def _extract_pdf(path: Path, ocr_language: str) -> str:
    import pymupdf

    parts: list[str] = []
    with pymupdf.open(path) as document:
        parts.extend(_metadata_section({
            key.capitalize(): value
            for key, value in (document.metadata or {}).items()
            if key not in {"format", "encryption"}
        }))
        for index, page in enumerate(document, start=1):
            parts.append(f"## Page {index}")
            text = page.get_text("text").strip()
            if not text:
                pixmap = page.get_pixmap(matrix=pymupdf.Matrix(2, 2), alpha=False)
                parts.append(_ocr_image_bytes(pixmap.tobytes("png"), ocr_language))
                continue

            # A selectable text layer does not cover text inside embedded images.
            # Render each image region so masks and placement are respected.
            seen_regions: set[tuple[float, ...]] = set()
            for image in page.get_image_info():
                region = (pymupdf.Rect(image["bbox"]) * page.rotation_matrix) & page.rect
                coordinates = tuple(region)
                if region.is_empty or coordinates in seen_regions:
                    continue
                seen_regions.add(coordinates)
                pixmap = page.get_pixmap(
                    matrix=pymupdf.Matrix(2, 2), clip=region, alpha=False
                )
                ocr_text = _ocr_image_bytes(pixmap.tobytes("png"), ocr_language)
                text = _append_ocr_lines(text, ocr_text)
            parts.append(text)
    return "\n\n".join(parts)


def _append_ocr_lines(text: str, ocr_text: str) -> str:
    """Append OCR lines without repeating exact lines in the text layer."""
    seen = {" ".join(line.split()) for line in text.splitlines()}
    extra = []
    for line in ocr_text.splitlines():
        key = " ".join(line.split())
        if key and key not in seen:
            seen.add(key)
            extra.append(line.strip())
    return "\n".join([text, *extra])


def _office_metadata(properties: Any) -> list[str]:
    """Read the shared DOCX/PPTX core text properties."""
    fields = (
        "author", "last_modified_by", "title", "subject", "keywords", "comments",
        "category", "identifier", "language", "content_status", "version",
    )
    return _metadata_section({
        name.replace("_", " ").capitalize(): getattr(properties, name, None)
        for name in fields
    })


def _metadata_section(values: dict[str, Any]) -> list[str]:
    """Keep metadata in the same Markdown stream as document content."""
    lines = [
        f"- {name}: {str(value).strip()}"
        for name, value in values.items()
        if value is not None and str(value).strip()
    ]
    return ["## Document metadata", "\n".join(lines)] if lines else []


def _extract_image(path: Path, ocr_language: str) -> str:
    from PIL import Image

    with Image.open(path) as image:
        text = _run_tesseract(image, ocr_language)
    return f"## OCR text\n\n{text.strip()}"


def _ocr_image_bytes(image_bytes: bytes, ocr_language: str) -> str:
    from PIL import Image

    with Image.open(BytesIO(image_bytes)) as image:
        return _run_tesseract(image, ocr_language).strip()


def _run_tesseract(image: Any, ocr_language: str) -> str:
    try:
        import pytesseract
    except ImportError as exc:
        raise ExtractionError(
            "OCR requires pytesseract and Pillow. Install the project requirements."
        ) from exc

    try:
        return pytesseract.image_to_string(image, lang=ocr_language)
    except pytesseract.TesseractNotFoundError as exc:
        raise ExtractionError(
            "OCR requires the Tesseract executable. Install Tesseract and add it to PATH."
        ) from exc
    except pytesseract.TesseractError as exc:
        raise ExtractionError(
            f"Tesseract could not use OCR language '{ocr_language}': {exc}"
        ) from exc


def _rows_to_markdown(rows: list[list[Any]], empty_message: str) -> str:
    if not rows:
        return empty_message

    width = max(len(row) for row in rows)
    padded_rows = [row + [""] * (width - len(row)) for row in rows]
    escaped_rows = [[_escape_cell(value) for value in row] for row in padded_rows]
    header = escaped_rows[0]
    separator = ["---"] * width

    lines = [_markdown_row(header), _markdown_row(separator)]
    lines.extend(_markdown_row(row) for row in escaped_rows[1:])
    return "\n".join(lines)


def _escape_cell(value: Any) -> str:
    return str(value).strip().replace("|", "\\|").replace("\r\n", "<br>").replace("\n", "<br>")


def _markdown_row(row: list[str]) -> str:
    return f"| {' | '.join(row)} |"
