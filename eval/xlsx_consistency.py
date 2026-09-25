from collections import defaultdict
from openpyxl import load_workbook


FIELD_TYPES = {
    "name": "PERSON",
    "email": "EMAIL",
    "phone": "PHONE",
    "company": "COMPANY",
    "iban": "IBAN",
    "date": "DATE",
    "personal_id": "PERSONAL_ID",
    "plate": "PLATE",
    "address": "ADDRESS",
}


def collect_original_occurrences(file_path):
    """
    Find every sensitive value in the original workbook
    and remember where it appears.
    """

    workbook = load_workbook(file_path, data_only=False)

    occurrences = []

    # Metadata
    if workbook.properties.creator:
        occurrences.append(
            {
                "type": "PERSON",
                "original": str(workbook.properties.creator),
                "location": "META:creator",
            }
        )

    if workbook.properties.lastModifiedBy:
        occurrences.append(
            {
                "type": "PERSON",
                "original": str(workbook.properties.lastModifiedBy),
                "location": "META:lastModifiedBy",
            }
        )

    # Worksheets including hidden sheets
    for sheet in workbook.worksheets:

        headers = {}

        for cell in sheet[1]:
            if cell.value is not None:
                headers[cell.column] = (
                    str(cell.value)
                    .strip()
                    .lower()
                )

        for row in sheet.iter_rows(min_row=2):

            for cell in row:

                if cell.value is None:
                    continue

                header = headers.get(cell.column)

                if header not in FIELD_TYPES:
                    continue

                occurrences.append(
                    {
                        "type": FIELD_TYPES[header],
                        "original": str(cell.value),
                        "location": f"{sheet.title}!{cell.coordinate}",
                    }
                )

    return occurrences


def get_output_value(workbook, location):
    """
    Read the anonymized value from exactly the same location.
    """

    if location == "META:creator":
        return str(workbook.properties.creator)

    if location == "META:lastModifiedBy":
        return str(workbook.properties.lastModifiedBy)

    sheet_name, cell_coordinate = location.rsplit("!", 1)

    sheet = workbook[sheet_name]

    value = sheet[cell_coordinate].value

    if value is None:
        return ""

    return str(value)


def evaluate_consistency(original_path, anonymized_path):

    occurrences = collect_original_occurrences(
        original_path
    )

    anonymized_workbook = load_workbook(
        anonymized_path,
        data_only=False
    )

    # Group identical original values together.
    groups = defaultdict(list)

    for item in occurrences:

        key = (
            item["type"],
            item["original"]
        )

        groups[key].append(
            item["location"]
        )

    # Only repeated values can test consistency.
    repeated_groups = {
        key: locations
        for key, locations in groups.items()
        if len(locations) > 1
    }

    consistent_groups = 0
    inconsistent_groups = 0

    print("=" * 70)
    print("XLSX CONSISTENCY EVALUATION")
    print("=" * 70)

    print(
        f"\nRepeated sensitive values found: "
        f"{len(repeated_groups)}"
    )

    for (
        entity_type,
        original_value
    ), locations in repeated_groups.items():

        outputs = []

        for location in locations:

            output_value = get_output_value(
                anonymized_workbook,
                location
            )

            outputs.append(
                (
                    location,
                    output_value
                )
            )

        unique_outputs = {
            value
            for location, value in outputs
        }

        is_consistent = (
            len(unique_outputs) == 1
        )

        print("\n" + "-" * 70)

        print(
            f"{entity_type}: "
            f"{original_value}"
        )

        print(
            f"Occurrences: {len(locations)}"
        )

        for location, output_value in outputs:

            print(
                f"  {location}"
                f" -> {output_value}"
            )

        if is_consistent:

            print("Result: CONSISTENT")
            consistent_groups += 1

        else:

            print("Result: INCONSISTENT")
            inconsistent_groups += 1

    total = len(repeated_groups)

    # -------------------------------------------------
    # SCORE
    # -------------------------------------------------

    if total == 0:

        print(
            "\nNot enough repeated sensitive values "
            "to evaluate consistency."
        )

        return None

    consistency_rate = (
        consistent_groups / total
    ) * 100

    if consistency_rate == 100:
        score = 4

    elif consistency_rate >= 75:
        score = 3

    elif consistency_rate >= 50:
        score = 2

    elif consistency_rate > 0:
        score = 1

    else:
        score = 0

    print("\n" + "=" * 70)
    print("[RESULT]")

    print(
        f"Consistent groups: "
        f"{consistent_groups}/{total}"
    )

    print(
        f"Consistency rate: "
        f"{consistency_rate:.1f}%"
    )

    print(
        f"Consistency Score: "
        f"{score}/4"
    )

    return score


if __name__ == "__main__":

    original_path = input(
        "Enter ORIGINAL XLSX file path: "
    ).strip().strip('"')

    anonymized_path = input(
        "Enter ANONYMIZED XLSX file path: "
    ).strip().strip('"')

    evaluate_consistency(
        original_path,
        anonymized_path
    )