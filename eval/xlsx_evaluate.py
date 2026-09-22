from pathlib import Path
from openpyxl import load_workbook


def collect_values(file_path):
    """
    Collect all cell values and important metadata
    from an XLSX file.
    """

    workbook = load_workbook(file_path, data_only=False)

    values = []

    # Metadata
    if workbook.properties.creator:
        values.append(str(workbook.properties.creator))

    if workbook.properties.lastModifiedBy:
        values.append(str(workbook.properties.lastModifiedBy))

    # All worksheets, including hidden sheets
    for sheet in workbook.worksheets:
        for row in sheet.iter_rows():
            for cell in row:
                if cell.value is not None:
                    values.append(str(cell.value))

    return values


def get_original_sensitive_values(file_path):
    """
    Read sensitive values from the original XLSX
    based on the column headers.
    """

    sensitive_fields = {
        "name",
        "email",
        "phone",
        "company",
        "iban",
        "date",
        "personal_id",
        "plate",
        "address",
    }

    workbook = load_workbook(file_path, data_only=False)

    sensitive_values = set()

    # Metadata
    if workbook.properties.creator:
        sensitive_values.add(str(workbook.properties.creator))

    if workbook.properties.lastModifiedBy:
        sensitive_values.add(str(workbook.properties.lastModifiedBy))

    # Worksheets
    for sheet in workbook.worksheets:

        headers = {}

        for cell in sheet[1]:
            if cell.value is not None:
                headers[cell.column] = str(cell.value).strip().lower()

        for row in sheet.iter_rows(min_row=2):

            for cell in row:

                if cell.value is None:
                    continue

                header = headers.get(cell.column)

                if header in sensitive_fields:
                    sensitive_values.add(str(cell.value))

    return sensitive_values


def evaluate_leakage(original_path, anonymized_path):

    original_sensitive = get_original_sensitive_values(original_path)
    anonymized_values = collect_values(anonymized_path)

    leaks = []

    # Search original sensitive values inside anonymized output
    for sensitive_value in original_sensitive:

        for output_value in anonymized_values:

            if sensitive_value.lower() in output_value.lower():
                leaks.append(sensitive_value)
                break

    # Remove duplicates
    leaks = sorted(set(leaks))

    print("=" * 70)
    print("XLSX LEAKAGE EVALUATION")
    print("=" * 70)

    print(f"\nOriginal sensitive values: {len(original_sensitive)}")
    print(f"Sensitive values found in output: {len(leaks)}")

    if leaks:
        print("\n[LEAKS FOUND]")

        for leak in leaks:
            print(f"- {leak}")
    else:
        print("\nNo original sensitive values found.")

    # Leakage score
    if len(leaks) == 0:
        score = 4

    elif len(leaks) == 1:
        score = 3

    elif len(leaks) <= 3:
        score = 2

    elif len(leaks) <= 5:
        score = 1

    else:
        score = 0

    print("\n[RESULT]")
    print(f"Data Leakage Score: {score}/4")

    if score >= 3:
        print("Technical leakage test: PASS")
    else:
        print("Technical leakage test: FAIL")

    return score, leaks


if __name__ == "__main__":

    original_path = input(
        "Enter ORIGINAL XLSX file path: "
    ).strip().strip('"')

    anonymized_path = input(
        "Enter ANONYMIZED XLSX file path: "
    ).strip().strip('"')

    evaluate_leakage(
        original_path,
        anonymized_path
    )