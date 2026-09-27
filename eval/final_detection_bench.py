#!/usr/bin/env python3

import collections
import importlib.metadata
import json
import os
from pathlib import Path

from openpyxl import load_workbook
from pptx import Presentation
import pymupdf


# ---------------------------------------------------------
# Presidio configuration
# Same mapping/configuration as eval/bench.py
# ---------------------------------------------------------

TYPE_MAP = {
    "PERSON": "PERSON",
    "EMAIL_ADDRESS": "EMAIL",
    "PHONE_NUMBER": "PHONE",
    "IBAN_CODE": "IBAN",
    "LOCATION": "ADDRESS",
    "ORGANIZATION": "COMPANY",
    "DATE_TIME": "DATE",
    "FI_PERSONAL_IDENTITY_CODE": "PERSONAL_ID",
}

MODELS = {
    "en": "en_core_web_sm",
    "fi": "fi_core_news_sm",
}

THRESHOLD = 0.0

FORMATS = {
    "xlsx",
    "pptx",
    "pdf",
}


# ---------------------------------------------------------
# Build Presidio analyzer
# ---------------------------------------------------------

def build_analyzer():
    from presidio_analyzer import AnalyzerEngine
    from presidio_analyzer.nlp_engine import NlpEngineProvider

    config = {
        "nlp_engine_name": "spacy",
        "models": [
            {
                "lang_code": code,
                "model_name": model,
            }
            for code, model in MODELS.items()
        ],
    }

    engine = NlpEngineProvider(
        nlp_configuration=config
    ).create_engine()

    return AnalyzerEngine(
        nlp_engine=engine,
        supported_languages=list(MODELS),
    )


# ---------------------------------------------------------
# Body extraction
# ---------------------------------------------------------

def extract_xlsx_body(path):
    """
    Extract only visible worksheet content.

    Hidden worksheets and metadata are intentionally
    excluded because the detection benchmark is body-only.
    """

    workbook = load_workbook(
        path,
        data_only=False,
    )

    parts = []

    for sheet in workbook.worksheets:

        if sheet.sheet_state != "visible":
            continue

        for row in sheet.iter_rows():

            values = []

            for cell in row:

                if cell.value is not None:
                    values.append(
                        str(cell.value)
                    )

            if values:
                parts.append(
                    " | ".join(values)
                )

    workbook.close()

    return "\n".join(parts)


def extract_shape_text(shape):
    texts = []

    if hasattr(shape, "text"):

        value = shape.text.strip()

        if value:
            texts.append(value)

    if getattr(shape, "has_table", False):

        for row in shape.table.rows:
            for cell in row.cells:

                value = cell.text.strip()

                if value:
                    texts.append(value)

    # Grouped shape
    if getattr(shape, "shape_type", None) == 6:

        for child in shape.shapes:
            texts.extend(
                extract_shape_text(child)
            )

    return texts


def extract_pptx_body(path):
    """
    Extract slide body only.

    Speaker notes and metadata are intentionally excluded.
    """

    presentation = Presentation(path)

    parts = []

    for slide in presentation.slides:

        for shape in slide.shapes:

            parts.extend(
                extract_shape_text(shape)
            )

    return "\n".join(parts)


def extract_pdf_body(path):
    """
    Extract native page text only.

    PDF metadata is intentionally excluded.
    """

    document = pymupdf.open(path)

    parts = []

    for page in document:

        text = page.get_text(
            "text"
        ).strip()

        if text:
            parts.append(text)

    document.close()

    return "\n".join(parts)


def extract_body(path, fmt):

    if fmt == "xlsx":
        return extract_xlsx_body(path)

    if fmt == "pptx":
        return extract_pptx_body(path)

    if fmt == "pdf":
        return extract_pdf_body(path)

    raise ValueError(
        f"Unsupported format: {fmt}"
    )


# ---------------------------------------------------------
# Same coverage logic as eval/bench.py
# ---------------------------------------------------------

def found(text, value, results):

    states = []

    start = text.find(value)

    while start >= 0:

        end = start + len(value)

        if any(
            result.start <= start
            and result.end >= end
            for result in results
        ):
            states.append("covered")

        elif any(
            result.start < end
            and result.end > start
            for result in results
        ):
            states.append("partial")

        else:
            states.append("")

        start = text.find(
            value,
            start + 1,
        )

    if states and all(
        state == "covered"
        for state in states
    ):
        return "covered"

    if any(states):
        return "partial"

    return ""


# ---------------------------------------------------------
# Benchmark
# ---------------------------------------------------------

def main(root):

    root = Path(root)

    labels_dir = (
        root / "labels"
    )

    analyzer = build_analyzer()

    # (format, language, type)
    # -> [planted, covered, partial, missed]
    counts = collections.defaultdict(
        lambda: [0, 0, 0, 0]
    )

    misses = []

    documents_processed = 0
    body_labels = 0

    label_files = sorted(
        labels_dir.glob("*.json")
    )

    for label_file in label_files:

        with open(
            label_file,
            encoding="utf-8",
        ) as file:
            data = json.load(file)

        fmt = data["format"]

        if fmt not in FORMATS:
            continue

        language = data["language"]

        document_path = (
            root / data["file"]
        )

        if not document_path.exists():

            print(
                f"[MISSING FILE] "
                f"{document_path}"
            )

            continue

        text = extract_body(
            document_path,
            fmt,
        )

        results = analyzer.analyze(
            text=text,
            language=language,
            score_threshold=THRESHOLD,
        )

        by_type = (
            collections.defaultdict(
                list
            )
        )

        for result in results:

            mapped_type = TYPE_MAP.get(
                result.entity_type
            )

            if mapped_type:
                by_type[
                    mapped_type
                ].append(result)

        documents_processed += 1

        # ---------------------------------------------
        # BODY labels only
        # ---------------------------------------------

        for entity in data["entities"]:

            if entity["location"] != "body":
                continue

            body_labels += 1

            entity_type = entity["type"]
            value = entity["value"]

            key = (
                fmt,
                language,
                entity_type,
            )

            counts[key][0] += 1

            result = found(
                text,
                value,
                by_type[entity_type],
            )

            if result == "covered":

                counts[key][1] += 1

            elif result == "partial":

                counts[key][2] += 1

                misses.append({
                    "file": data["file"],
                    "language": language,
                    "format": fmt,
                    "type": entity_type,
                    "value": value,
                    "detected": "partial",
                })

            else:

                counts[key][3] += 1

                misses.append({
                    "file": data["file"],
                    "language": language,
                    "format": fmt,
                    "type": entity_type,
                    "value": value,
                    "detected": "nothing",
                })

    # -------------------------------------------------
    # PRINT RESULTS
    # -------------------------------------------------

    print()
    print("=" * 86)
    print("FINAL PRESIDIO DETECTION BENCHMARK")
    print("=" * 86)

    print(
        f"\nDocuments processed: "
        f"{documents_processed}"
    )

    print(
        f"Body ground-truth labels: "
        f"{body_labels}"
    )

    print()

    header = (
        f"{'FORMAT':<7}"
        f"{'LANG':<6}"
        f"{'TYPE':<14}"
        f"{'TOTAL':>7}"
        f"{'FULL':>7}"
        f"{'PART':>7}"
        f"{'MISS':>7}"
        f"{'RECALL':>10}"
    )

    print(header)
    print("-" * len(header))

    for key in sorted(counts):

        fmt, language, entity_type = key

        planted, covered, partial, missed = (
            counts[key]
        )

        recall = (
            covered / planted
            if planted
            else 0
        )

        print(
            f"{fmt.upper():<7}"
            f"{language.upper():<6}"
            f"{entity_type:<14}"
            f"{planted:>7}"
            f"{covered:>7}"
            f"{partial:>7}"
            f"{missed:>7}"
            f"{recall:>9.1%}"
        )

    # -------------------------------------------------
    # FORMAT + LANGUAGE totals
    # -------------------------------------------------

    print()
    print("=" * 86)
    print("TOTALS BY FORMAT AND LANGUAGE")
    print("=" * 86)

    totals = collections.defaultdict(
        lambda: [0, 0, 0, 0]
    )

    for (
        fmt,
        language,
        entity_type,
    ), values in counts.items():

        for index in range(4):
            totals[
                (fmt, language)
            ][index] += values[index]

    for key in sorted(totals):

        fmt, language = key

        planted, covered, partial, missed = (
            totals[key]
        )

        recall = (
            covered / planted
            if planted
            else 0
        )

        print(
            f"{fmt.upper():<6} "
            f"{language.upper():<2} | "
            f"Full {covered}/{planted} "
            f"({recall:.2%}) | "
            f"Partial {partial} | "
            f"Missed {missed}"
        )

    # -------------------------------------------------
    # Combined totals
    # -------------------------------------------------

    planted = sum(
        value[0]
        for value in counts.values()
    )

    covered = sum(
        value[1]
        for value in counts.values()
    )

    partial = sum(
        value[2]
        for value in counts.values()
    )

    missed = sum(
        value[3]
        for value in counts.values()
    )

    print()
    print("=" * 86)
    print("COMBINED")
    print("=" * 86)

    print(
        f"Ground truth: {planted}"
    )

    print(
        f"Fully detected: "
        f"{covered}/{planted} "
        f"({covered / planted:.2%})"
    )

    print(
        f"Partial: "
        f"{partial}/{planted} "
        f"({partial / planted:.2%})"
    )

    print(
        f"Missed: "
        f"{missed}/{planted} "
        f"({missed / planted:.2%})"
    )

    # -------------------------------------------------
    # Save results
    # -------------------------------------------------

    output_dir = Path(
        "eval/results"
    )

    output_dir.mkdir(
        parents=True,
        exist_ok=True,
    )

    model_versions = {
        package:
            importlib.metadata.version(
                package
            )
        for package in (
            "presidio-analyzer",
            "spacy",
        )
    }

    for model in MODELS.values():

        model_versions[model] = (
            importlib.metadata.version(
                model
            )
        )

    output = {
        "detector": "presidio-analyzer",
        "versions": model_versions,
        "models": MODELS,
        "score_threshold": THRESHOLD,
        "scope": (
            "Final detection benchmark for XLSX, PPTX and PDF. "
            "Body labels only. Metadata, notes and hidden sheets "
            "belong to extraction evaluation."
        ),
        "documents_processed": (
            documents_processed
        ),
        "body_labels": body_labels,
        "counts": {},
    }

    for key, values in sorted(
        counts.items()
    ):

        fmt, language, entity_type = key

        planted, covered, partial, missed = (
            values
        )

        output["counts"][
            f"{fmt}/{language}/{entity_type}"
        ] = {
            "planted": planted,
            "covered": covered,
            "partial": partial,
            "missed": missed,
        }

    with open(
        output_dir
        / "final-presidio.json",
        "w",
        encoding="utf-8",
    ) as file:

        json.dump(
            output,
            file,
            indent=2,
            ensure_ascii=False,
        )

    with open(
        output_dir
        / "final-misses.json",
        "w",
        encoding="utf-8",
    ) as file:

        json.dump(
            misses,
            file,
            indent=2,
            ensure_ascii=False,
        )

    print()
    print(
        "Saved:"
    )

    print(
        output_dir
        / "final-presidio.json"
    )

    print(
        output_dir
        / "final-misses.json"
    )


if __name__ == "__main__":

    import sys

    root = (
        sys.argv[1]
        if len(sys.argv) > 1
        else "bench_final_detection"
    )

    main(root)

