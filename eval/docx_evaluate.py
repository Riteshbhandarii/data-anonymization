import docx 
from presidio_analyzer import AnalyzerEngine

def evaluate_docx_leakage(orig_path, anon_path):
    orig_doc = docx.Document(orig_path)
    anon_doc = docx.Document(anon_path)

    orig_text = "\n".join([p.text for p in orig_doc.paragraphs if p.text.strip()])
    anon_text = "\n".join([p.text for p in anon_doc.paragraphs if p.text.strip()])

    analyzer = AnalyzerEngine()
    results = analyzer.analyze(text=orig_text, language="en")

    original_sensitive = set([orig_text[r.start:r.end] for r in results if len(orig_text[r.start:r.end]) > 3])

    leaks = []
    for sensitive_value in original_sensitive:
        if sensitive_value.lower() in anon_text.lower():
            leaks.append(sensitive_value)

    print("=" * 70)
    print("DOCX LEAKAGE EVALUATION")
    print("=" * 70)
    print(f"\nOriginal sensitive value extracted: {len(original_sensitive)}")
    print(f"Sensitive values found in output: {len(leaks)}")

    if leaks:
        print("\n[LEAK FOUND]")
        for leak in leaks: print(f"- {leak}")
    else:
        print("\nNo original sensitive values found in output.")

    score = 4 if len(leaks) == 0 else (3 if len(leaks) == 1 else (2 if len(leaks) <= 3 else (1 if len(leaks) <= 5 else 0)))

    print("\n[RESULT]")
    print(f"Data Leakage Score: {score}/4")
    print("Technical leakage test: PASS" if score >= 3 else "Technical leakage test: FAIL")

if __name__ == "__main__":
    orig_path = input("Enter ORIGINAL DOCX file path: ").strip().strip('"')
    anon_path = input("Enter ANONYMIZED DOCX file path: ").strip().strip('"')
    evaluate_docx_leakage(orig_path, anon_path)