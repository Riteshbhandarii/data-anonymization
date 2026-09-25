from datetime import datetime, date
from collections import defaultdict
from openpyxl import load_workbook


FIELD_TYPES = {
    "name": "PERSON",
    "email": "EMAIL",
    "phone": "PHONE",
    "company": "COMPANY",
    "iban": "IBAN",
    "date": "DATE",
    "personal_id": "PERSONAL_ID",
    "plate": "PLATE",
    "address": "ADDRESS",
}


def get_headers(sheet):
    headers = {}

    for cell in sheet[1]:
        if cell.value is not None:
            headers[cell.column] = str(cell.value).strip()

    return headers


def collect_sensitive_cells(file_path):
    """
    Collect sensitive cells and their locations.
    """

    workbook = load_workbook(file_path, data_only=False)

    items = []

    for sheet in workbook.worksheets:

        headers = {}

        for cell in sheet[1]:
            if cell.value is not None:
                headers[cell.column] = (
                    str(cell.value)
                    .strip()
                    .lower()
                )

        for row in sheet.iter_rows(min_row=2):

            for cell in row:

                if cell.value is None:
                    continue

                header = headers.get(cell.column)

                if header not in FIELD_TYPES:
                    continue

                items.append(
                    {
                        "type": FIELD_TYPES[header],
                        "original": str(cell.value),
                        "sheet": sheet.title,
                        "cell": cell.coordinate,
                    }
                )

    # Metadata
    if workbook.properties.creator:
        items.append(
            {
                "type": "PERSON",
                "original": str(workbook.properties.creator),
                "sheet": "__METADATA__",
                "cell": "creator",
            }
        )

    if workbook.properties.lastModifiedBy:
        items.append(
            {
                "type": "PERSON",
                "original": str(workbook.properties.lastModifiedBy),
                "sheet": "__METADATA__",
                "cell": "lastModifiedBy",
            }
        )

    return items


def get_output_value(workbook, item):

    if item["sheet"] == "__METADATA__":

        if item["cell"] == "creator":
            return str(
                workbook.properties.creator or ""
            )

        if item["cell"] == "lastModifiedBy":
            return str(
                workbook.properties.lastModifiedBy or ""
            )

    value = workbook[
        item["sheet"]
    ][
        item["cell"]
    ].value

    if value is None:
        return ""

    return str(value)


# ---------------------------------------------------------
# CHECK 1: WORKBOOK STRUCTURE
# ---------------------------------------------------------

def check_structure(original_path, anonymized_path):

    original = load_workbook(
        original_path,
        data_only=False
    )

    anonymized = load_workbook(
        anonymized_path,
        data_only=False
    )

    if original.sheetnames != anonymized.sheetnames:
        return False

    for sheet_name in original.sheetnames:

        original_sheet = original[sheet_name]
        anonymized_sheet = anonymized[sheet_name]

        # Hidden / visible state
        if (
            original_sheet.sheet_state
            != anonymized_sheet.sheet_state
        ):
            return False

        # Headers
        if (
            get_headers(original_sheet)
            != get_headers(anonymized_sheet)
        ):
            return False

    return True


# ---------------------------------------------------------
# CHECK 2: ROW / COLUMN ASSOCIATIONS
# ---------------------------------------------------------

def check_cell_positions(original_path, anonymized_path):

    original_items = collect_sensitive_cells(
        original_path
    )

    anonymized = load_workbook(
        anonymized_path,
        data_only=False
    )

    for item in original_items:

        output = get_output_value(
            anonymized,
            item
        )

        # Information should still occupy
        # the corresponding logical position.
        if output == "":
            return False

    return True


# ---------------------------------------------------------
# CHECK 3: ENTITY DISTINGUISHABILITY
# ---------------------------------------------------------

def check_distinguishability(
    original_path,
    anonymized_path
):

    original_items = collect_sensitive_cells(
        original_path
    )

    anonymized = load_workbook(
        anonymized_path,
        data_only=False
    )

    mappings = defaultdict(
        lambda: defaultdict(set)
    )

    for item in original_items:

        output = get_output_value(
            anonymized,
            item
        )

        mappings[
            item["type"]
        ][
            item["original"]
        ].add(output)

    checked_types = 0

    for entity_type, original_mapping in mappings.items():

        # Need at least 2 different original values
        # to test whether distinction was preserved.
        if len(original_mapping) < 2:
            continue

        checked_types += 1

        representative_outputs = []

        for original, outputs in original_mapping.items():

            # Same original should not produce
            # multiple different outputs.
            if len(outputs) != 1:
                return False

            representative_outputs.append(
                next(iter(outputs))
            )

        # Different originals should remain
        # distinguishable.
        if (
            len(set(representative_outputs))
            != len(representative_outputs)
        ):
            return False

    # No comparable types means we cannot prove
    # distinguishability.
    if checked_types == 0:
        return False

    return True


# ---------------------------------------------------------
# CHECK 4: TEMPORAL RELATIONSHIP
# ---------------------------------------------------------

def parse_context_date(value):
    """
    Parse exact dates or generalized years.

    Supports:
    1986-01-16
    1986-01-16 00:00:00
    1986-**-**
    1986
    """

    if isinstance(value, datetime):
        return value

    if isinstance(value, date):
        return datetime.combine(
            value,
            datetime.min.time()
        )

    value = str(value)

    # Exact date / datetime
    formats = [
        "%Y-%m-%d",
        "%Y-%m-%d %H:%M:%S",
    ]

    for fmt in formats:

        try:
            return datetime.strptime(
                value,
                fmt
            )

        except ValueError:
            pass

    # Generalized or masked date:
    # use the year for ordering.
    if (
        len(value) >= 4
        and value[:4].isdigit()
    ):

        return datetime(
            int(value[:4]),
            1,
            1
        )

    return None


def check_temporal_relationship(
    original_path,
    anonymized_path
):

    original_items = collect_sensitive_cells(
        original_path
    )

    anonymized = load_workbook(
        anonymized_path,
        data_only=False
    )

    date_pairs = []

    for item in original_items:

        if item["type"] != "DATE":
            continue

        original_date = parse_context_date(
            item["original"]
        )

        output_value = get_output_value(
            anonymized,
            item
        )

        anonymized_date = parse_context_date(
            output_value
        )

        if (
            original_date is None
            or anonymized_date is None
        ):
            return False

        date_pairs.append(
            (
                original_date,
                anonymized_date
            )
        )

    if len(date_pairs) < 2:
        return False

    # Compare ordering of every date pair.
    for i in range(len(date_pairs)):

        for j in range(
            i + 1,
            len(date_pairs)
        ):

            original_a = date_pairs[i][0]
            original_b = date_pairs[j][0]

            anonymized_a = date_pairs[i][1]
            anonymized_b = date_pairs[j][1]

            original_order = (
                original_a < original_b
            )

            anonymized_order = (
                anonymized_a < anonymized_b
            )

            if original_order != anonymized_order:
                return False

    return True


# ---------------------------------------------------------
# MAIN EVALUATION
# ---------------------------------------------------------

def evaluate_context(
    original_path,
    anonymized_path
):

    print("=" * 70)
    print("XLSX CONTEXT PRESERVATION EVALUATION")
    print("=" * 70)

    results = {
        "Workbook structure": check_structure(
            original_path,
            anonymized_path
        ),

        "Row/column associations": check_cell_positions(
            original_path,
            anonymized_path
        ),

        "Entity distinguishability":
            check_distinguishability(
                original_path,
                anonymized_path
            ),

        "Temporal relationship":
            check_temporal_relationship(
                original_path,
                anonymized_path
            ),
    }

    print()

    for name, passed in results.items():

        status = (
            "PASS"
            if passed
            else "FAIL"
        )

        print(
            f"{name:<30} {status}"
        )

    score = sum(
        1
        for passed in results.values()
        if passed
    )

    print("\n[RESULT]")

    print(
        f"Context Preservation Score: "
        f"{score}/4"
    )

    if score == 4:
        meaning = (
            "Context is almost completely preserved"
        )

    elif score == 3:
        meaning = (
            "Most context remains"
        )

    elif score == 2:
        meaning = (
            "Basic meaning remains"
        )

    elif score == 1:
        meaning = (
            "Important relationships are lost"
        )

    else:
        meaning = (
            "Meaning of the data is destroyed"
        )

    print(
        f"Interpretation: {meaning}"
    )

    return score, results


if __name__ == "__main__":

    original_path = input(
        "Enter ORIGINAL XLSX file path: "
    ).strip().strip('"')

    anonymized_path = input(
        "Enter ANONYMIZED XLSX file path: "
    ).strip().strip('"')

    evaluate_context(
        original_path,
        anonymized_path
    )