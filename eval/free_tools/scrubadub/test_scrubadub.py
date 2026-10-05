from pathlib import Path
import csv

import scrubadub
from docx import Document
from pypdf import PdfReader
from pptx import Presentation
from openpyxl import load_workbook


# ============================================================
# 1. EXTRACT TEXT FROM EACH FILE FORMAT
# ============================================================

def extract_csv(path):
    """Extract text from a CSV file."""
    lines = []

    with open(path, "r", encoding="utf-8-sig", errors="replace") as f:
        reader = csv.reader(f)

        for row in reader:
            lines.append(" | ".join(str(value) for value in row))

    return "\n".join(lines)


def extract_docx(path):
    """Extract paragraphs and table contents from a DOCX file."""
    document = Document(path)
    lines = []

    for paragraph in document.paragraphs:
        if paragraph.text.strip():
            lines.append(paragraph.text)

    for table in document.tables:
        for row in table.rows:
            values = [
                cell.text.strip()
                for cell in row.cells
                if cell.text.strip()
            ]

            if values:
                lines.append(" | ".join(values))

    return "\n".join(lines)


def extract_pdf(path):
    """Extract text from a text-based PDF."""
    reader = PdfReader(path)
    lines = []

    for page_number, page in enumerate(reader.pages, start=1):
        text = page.extract_text() or ""

        if text.strip():
            lines.append(f"[PAGE {page_number}]")
            lines.append(text)

    return "\n".join(lines)


def extract_pptx(path):
    """
    Extract visible slide text and speaker notes from PPTX.

    Speaker notes are important because sensitive information
    may be hidden there.
    """
    presentation = Presentation(path)
    lines = []

    for slide_number, slide in enumerate(presentation.slides, start=1):

        lines.append(f"[SLIDE {slide_number}]")

        for shape in slide.shapes:
            if hasattr(shape, "text"):
                text = shape.text.strip()

                if text:
                    lines.append(text)

        try:
            notes_slide = slide.notes_slide

            note_text = []

            for shape in notes_slide.shapes:
                if hasattr(shape, "text"):
                    text = shape.text.strip()

                    if text:
                        note_text.append(text)

            if note_text:
                lines.append("[SPEAKER NOTES]")
                lines.extend(note_text)

        except Exception:
            pass

    return "\n".join(lines)


def extract_xlsx(path):
    """Extract cell values from all worksheets in an XLSX file."""
    workbook = load_workbook(path, data_only=True)
    lines = []

    for sheet in workbook.worksheets:

        lines.append(f"[SHEET: {sheet.title}]")

        for row in sheet.iter_rows(values_only=True):

            values = [
                str(value)
                for value in row
                if value is not None
            ]

            if values:
                lines.append(" | ".join(values))

    workbook.close()

    return "\n".join(lines)


# ============================================================
# 2. CHOOSE THE CORRECT EXTRACTOR
# ============================================================

def extract_text(path):

    extension = path.suffix.lower()

    if extension == ".csv":
        return extract_csv(path)

    elif extension == ".docx":
        return extract_docx(path)

    elif extension == ".pdf":
        return extract_pdf(path)

    elif extension == ".pptx":
        return extract_pptx(path)

    elif extension == ".xlsx":
        return extract_xlsx(path)

    else:
        raise ValueError(f"Unsupported file type: {extension}")


# ============================================================
# 3. FIND ORIGINAL TEST FILES
# ============================================================

extensions = {
    "csv": ".csv",
    "docx": ".docx",
    "pdf": ".pdf",
    "pptx": ".pptx",
    "xlsx": ".xlsx",
}

files = []

for folder, extension in extensions.items():

    folder_path = Path(folder)

    if not folder_path.exists():
        print(f"WARNING: Folder not found: {folder}")
        continue

    for path in folder_path.glob(f"*{extension}"):

        filename = path.name.lower()

        if filename.startswith("anonymized_"):
            continue

        if not (
            filename.startswith("en_")
            or filename.startswith("fi_")
        ):
            continue

        files.append(path)


files = sorted(files)


# ============================================================
# 4. DISPLAY TEST SET INFORMATION
# ============================================================

print("=" * 70)
print("SCRUBADUB TEST")
print("=" * 70)

print(f"\nFound {len(files)} original test files.\n")


counts = {}

for path in files:

    file_format = path.suffix.lower().replace(".", "")

    if path.name.lower().startswith("en_"):
        language = "EN"
    elif path.name.lower().startswith("fi_"):
        language = "FI"
    else:
        language = "UNKNOWN"

    key = (file_format, language)

    counts[key] = counts.get(key, 0) + 1


print("FILES BY FORMAT/LANGUAGE:")

for key in sorted(counts):
    print(f"{key[0].upper()} {key[1]}: {counts[key]}")


# ============================================================
# 5. CREATE OUTPUT FOLDER
# ============================================================

output_folder = Path("scrubadub_results")
output_folder.mkdir(exist_ok=True)


# ============================================================
# 6. CREATE SUMMARY CSV
# ============================================================

summary_path = output_folder / "scrubadub_summary.csv"

summary_file = open(
    summary_path,
    "w",
    newline="",
    encoding="utf-8"
)

summary_writer = csv.writer(summary_file)

summary_writer.writerow([
    "filename",
    "format",
    "language",
    "status",
    "original_characters",
    "anonymized_characters",
    "changed",
    "error"
])


# ============================================================
# 7. PROCESS EVERY FILE
# ============================================================

success_count = 0
error_count = 0
changed_count = 0


for number, path in enumerate(files, start=1):

    print(
        f"[{number}/{len(files)}] "
        f"Testing {path}"
    )

    file_format = path.suffix.lower().replace(".", "")

    if path.name.lower().startswith("en_"):
        language = "EN"

    elif path.name.lower().startswith("fi_"):
        language = "FI"

    else:
        language = "UNKNOWN"

    try:

        # Extract original text
        original = extract_text(path)

        # Run scrubadub
        anonymized = scrubadub.clean(original)

        # Determine whether scrubadub changed anything
        changed = original != anonymized

        if changed:
            changed_count += 1

        # ----------------------------------------------------
        # Save detailed result
        # ----------------------------------------------------

        result_name = (
            f"{path.stem}_scrubadub_result.txt"
        )

        result_path = output_folder / result_name

        with open(
            result_path,
            "w",
            encoding="utf-8"
        ) as result_file:

            result_file.write(
                f"FILE: {path}\n"
            )

            result_file.write(
                f"FORMAT: {file_format.upper()}\n"
            )

            result_file.write(
                f"LANGUAGE: {language}\n"
            )

            result_file.write(
                f"CHANGED BY SCRUBADUB: {changed}\n"
            )

            result_file.write(
                "\n"
                + "=" * 70
                + "\nORIGINAL\n"
                + "=" * 70
                + "\n\n"
            )

            result_file.write(original)

            result_file.write(
                "\n\n"
                + "=" * 70
                + "\nANONYMIZED BY SCRUBADUB\n"
                + "=" * 70
                + "\n\n"
            )

            result_file.write(anonymized)

        # ----------------------------------------------------
        # Add summary row
        # ----------------------------------------------------

        summary_writer.writerow([
            path.name,
            file_format,
            language,
            "SUCCESS",
            len(original),
            len(anonymized),
            changed,
            ""
        ])

        success_count += 1

    except Exception as error:

        print(f"    ERROR: {error}")

        summary_writer.writerow([
            path.name,
            file_format,
            language,
            "ERROR",
            "",
            "",
            "",
            str(error)
        ])

        error_count += 1


# ============================================================
# 8. FINISH
# ============================================================

summary_file.close()


print("\n" + "=" * 70)
print("TEST COMPLETE")
print("=" * 70)

print(f"Files tested successfully: {success_count}")
print(f"Files with errors:        {error_count}")
print(f"Files changed:            {changed_count}")

print("\nResults saved in:")
print(output_folder.resolve())

print("\nSummary:")
print(summary_path.resolve())