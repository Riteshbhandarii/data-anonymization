import csv
import logging
from pathlib import Path

from datafog_core import scan
from docx import Document
from openpyxl import load_workbook
from pptx import Presentation
from pypdf import PdfReader

logger = logging.getLogger(__name__)

# ============================================================
# CONFIGURATION
# ============================================================

BASE_DIR = Path(__file__).resolve().parent

OUTPUT_DIR = BASE_DIR / "datafog_results"
OUTPUT_DIR.mkdir(exist_ok=True)


# ============================================================
# TEXT EXTRACTION
# ============================================================

def extract_csv(path):
    lines = []

    with open(
        path,
        "r",
        encoding="utf-8-sig",
        errors="replace"
    ) as f:

        reader = csv.reader(f)

        for row in reader:
            lines.append(
                " | ".join(str(value) for value in row)
            )

    return "\n".join(lines)


def extract_docx(path):
    document = Document(path)
    lines = []

    # Paragraphs
    for paragraph in document.paragraphs:

        if paragraph.text.strip():
            lines.append(paragraph.text)

    # Tables
    for table in document.tables:

        for row in table.rows:

            values = [
                cell.text.strip()
                for cell in row.cells
                if cell.text.strip()
            ]

            if values:
                lines.append(" | ".join(values))

    return "\n".join(lines)


def extract_pdf(path):
    reader = PdfReader(path)

    lines = []

    for page_number, page in enumerate(
        reader.pages,
        start=1
    ):

        text = page.extract_text() or ""

        if text.strip():

            lines.append(
                f"[PAGE {page_number}]"
            )

            lines.append(text)

    return "\n".join(lines)


def extract_pptx(path):
    presentation = Presentation(path)

    lines = []

    for slide_number, slide in enumerate(
        presentation.slides,
        start=1
    ):

        lines.append(
            f"[SLIDE {slide_number}]"
        )

        # Visible slide text
        for shape in slide.shapes:

            if hasattr(shape, "text"):

                text = shape.text.strip()

                if text:
                    lines.append(text)

       
        if slide.has_notes_slide:

            notes_slide = slide.notes_slide

            notes = []

            for shape in notes_slide.shapes:

                if hasattr(shape, "text"):

                    text = shape.text.strip()

                    if text:
                        notes.append(text)

            if notes:

                lines.append("[SPEAKER NOTES]")
                lines.extend(notes)


    return "\n".join(lines)


def extract_xlsx(path):
    workbook = load_workbook(
        path,
        data_only=True
    )

    lines = []

    for sheet in workbook.worksheets:

        lines.append(
            f"[SHEET: {sheet.title}]"
        )

        for row in sheet.iter_rows(
            values_only=True
        ):

            values = [
                str(value)
                for value in row
                if value is not None
            ]

            if values:
                lines.append(
                    " | ".join(values)
                )

    workbook.close()

    return "\n".join(lines)


def extract_text(path):

    extension = path.suffix.lower()

    if extension == ".csv":
        return extract_csv(path)

    elif extension == ".docx":
        return extract_docx(path)

    elif extension == ".pdf":
        return extract_pdf(path)

    elif extension == ".pptx":
        return extract_pptx(path)

    elif extension == ".xlsx":
        return extract_xlsx(path)

    else:
        raise ValueError(
            f"Unsupported file: {path}"
        )


# ============================================================
# DATAFOG REDACTION
# ============================================================

def redact_with_datafog(text):
    findings = scan(text)
    anonymized_text = text

    # Build replacements from DataFog findings
    replacements = []

    for finding in findings:
        entity_type = finding.entity_type
        matched_text = finding.matched_text

        if matched_text:
            replacements.append((matched_text, entity_type))

    replacements.sort(key=lambda item: len(item[0]), reverse=True)

    for matched_text, entity_type in replacements:
        anonymized_text = anonymized_text.replace(
            matched_text,
            f"[{entity_type}]"
        )

    return anonymized_text, findings

# ============================================================
# FIND ORIGINAL TEST FILES
# ============================================================

extensions = {
    "csv": ".csv",
    "docx": ".docx",
    "pdf": ".pdf",
    "pptx": ".pptx",
    "xlsx": ".xlsx",
}

files = []


for folder, extension in extensions.items():

    folder_path = BASE_DIR / folder

    if not folder_path.exists():

        print(
            f"WARNING: Folder not found: "
            f"{folder_path}"
        )

        continue


    for path in folder_path.glob(
        f"*{extension}"
    ):

        filename = path.name.lower()

        if filename.startswith(
            "anonymized_"
        ):
            continue

        if not filename.startswith(("en_", "fi_")):
            continue

        files.append(path)


files = sorted(files)


# ============================================================
# DISPLAY DATASET COUNTS
# ============================================================

print("=" * 70)
print("DATAFOG TEST")
print("=" * 70)

print(
    f"\nFound {len(files)} "
    f"original test files.\n"
)


counts = {}

for path in files:

    file_format = (
        path.suffix
        .lower()
        .replace(".", "")
    )

    if path.name.lower().startswith(
        "en_"
    ):
        language = "EN"

    elif path.name.lower().startswith(
        "fi_"
    ):
        language = "FI"

    else:
        language = "UNKNOWN"

    key = (
        file_format,
        language
    )

    counts[key] = (
        counts.get(key, 0) + 1
    )


print("FILES BY FORMAT/LANGUAGE:")

for key in sorted(counts):

    print(
        f"{key[0].upper()} "
        f"{key[1]}: "
        f"{counts[key]}"
    )


# ============================================================
# SUMMARY CSV
# ============================================================

summary_path = (
    OUTPUT_DIR /
    "datafog_summary.csv"
)


with open(
    summary_path,
    "w",
    newline="",
    encoding="utf-8-sig"
) as summary_file:


    summary_writer = csv.writer(
        summary_file
    )


    summary_writer.writerow([
        "filename",
        "format",
        "language",
        "status",
        "detections",
        "original_characters",
        "anonymized_characters",
        "changed",
        "error"
    ])


    # ============================================================
    # PROCESS FILES
    # ============================================================

    success_count = 0
    error_count = 0
    changed_count = 0
    total_detections = 0


    for number, path in enumerate(
        files,
        start=1
    ):

        print(
            f"[{number}/{len(files)}] "
            f"Testing {path.name}"
        )


        file_format = (
            path.suffix
            .lower()
            .replace(".", "")
        )


        if path.name.lower().startswith(
            "en_"
        ):
            language = "EN"

        else:
            language = "FI"


        try:

            # Extract document text
            original = extract_text(path)


            # Run DataFog
            anonymized, findings = (
                redact_with_datafog(original)
            )


            detection_count = len(findings)

            total_detections += (
                detection_count
            )


            changed = (
                original != anonymized
            )


            if changed:
                changed_count += 1


            # ====================================================
            # SAVE DETAILED RESULT
            # ====================================================

            result_name = (
                f"{path.stem}"
                f"_datafog_result.txt"
            )


            result_path = (
                OUTPUT_DIR /
                result_name
            )


            with open(
                result_path,
                "w",
                encoding="utf-8"
            ) as result_file:

                result_file.write(
                    f"FILE: {path}\n"
                )

                result_file.write(
                    f"FORMAT: "
                    f"{file_format.upper()}\n"
                )

                result_file.write(
                    f"LANGUAGE: "
                    f"{language}\n"
                )

                result_file.write(
                    f"DETECTIONS: "
                    f"{detection_count}\n"
                )

                result_file.write(
                    f"CHANGED BY DATAFOG: "
                    f"{changed}\n"
                )


                result_file.write(
                    "\n"
                    + "=" * 70
                    + "\nORIGINAL\n"
                    + "=" * 70
                    + "\n\n"
                )

                result_file.write(
                    original
                )


                result_file.write(
                    "\n\n"
                    + "=" * 70
                    + "\nANONYMIZED BY DATAFOG\n"
                    + "=" * 70
                    + "\n\n"
                )

                result_file.write(
                    anonymized
                )


                result_file.write(
                    "\n\n"
                    + "=" * 70
                    + "\nDATAFOG FINDINGS\n"
                    + "=" * 70
                    + "\n\n"
                )

                result_file.write(
                    repr(findings)
                )


            # ====================================================
            # SUMMARY
            # ====================================================

            summary_writer.writerow([
                path.name,
                file_format,
                language,
                "SUCCESS",
                detection_count,
                len(original),
                len(anonymized),
                changed,
                ""
            ])


            success_count += 1


        except Exception as error:
            logger.exception("Failed to process %s", path)

            print(
                f"    ERROR: {error}"
            )


            summary_writer.writerow([
                path.name,
                file_format,
                language,
                "ERROR",
                "",
                "",
                "",
                "",
                str(error)
            ])


            error_count += 1





# ============================================================
# FINAL OUTPUT
# ============================================================

print(
    "\n" + "=" * 70
)

print("DATAFOG TEST COMPLETE")

print("=" * 70)


print(
    f"Files tested successfully: "
    f"{success_count}"
)

print(
    f"Files with errors:        "
    f"{error_count}"
)

print(
    f"Files changed:            "
    f"{changed_count}"
)

print(
    f"Total detections:         "
    f"{total_detections}"
)


print("\nResults saved in:")

print(
    OUTPUT_DIR.resolve()
)


print("\nSummary:")

print(
    summary_path.resolve()
)

if not files or error_count:
    raise SystemExit(1)
