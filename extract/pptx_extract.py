from pathlib import Path
from pptx import Presentation


def extract_shape_text(shape):
    """
    Extract text from normal text shapes and tables.
    """

    texts = []

    # Normal text box / placeholder
    if hasattr(shape, "text") and shape.text:
        text = shape.text.strip()
        if text:
            texts.append(text)

    # Table
    if getattr(shape, "has_table", False):
        for row in shape.table.rows:
            for cell in row.cells:
                text = cell.text.strip()
                if text:
                    texts.append(text)

    # Grouped shapes
    if shape.shape_type == 6:  # GROUP
        for child in shape.shapes:
            texts.extend(
                extract_shape_text(child)
            )

    return texts


def extract_pptx(file_path):
    file_path = Path(file_path)

    prs = Presentation(file_path)

    print("=" * 70)
    print(f"FILE: {file_path.name}")
    print("=" * 70)

    # --------------------------------------------------
    # METADATA
    # --------------------------------------------------

    props = prs.core_properties

    print("\n[METADATA]")

    print(
        f"Creator: {props.author}"
    )

    print(
        f"Last modified by: "
        f"{props.last_modified_by}"
    )

    print(
        f"Title: {props.title}"
    )

    print(
        f"Subject: {props.subject}"
    )

    print(
        f"Keywords: {props.keywords}"
    )

    print(
        f"Comments: {props.comments}"
    )

    # --------------------------------------------------
    # SLIDES
    # --------------------------------------------------

    print("\n[SLIDES]")

    for slide_number, slide in enumerate(
        prs.slides,
        start=1,
    ):

        print(
            f"\nSlide {slide_number}"
        )

        found_text = False

        for shape in slide.shapes:

            values = extract_shape_text(
                shape
            )

            for value in values:

                print(
                    f"BODY | {value}"
                )

                found_text = True

        if not found_text:
            print(
                "BODY | [No text]"
            )

        # ----------------------------------------------
        # SPEAKER NOTES
        # ----------------------------------------------

        if slide.has_notes_slide:

            notes_slide = (
                slide.notes_slide
            )

            notes_frame = (
                notes_slide.notes_text_frame
            )

            if notes_frame is not None:

                notes_text = (
                    notes_frame.text.strip()
                )

                if notes_text:

                    print(
                        f"NOTES | {notes_text}"
                    )


if __name__ == "__main__":

    file_path = input(
        "Enter PPTX file path: "
    ).strip().strip('"')

    extract_pptx(
        file_path
    )