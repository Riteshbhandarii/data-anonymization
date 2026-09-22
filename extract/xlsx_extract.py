from pathlib import Path
from openpyxl import load_workbook


def inspect_xlsx(file_path):
    file_path = Path(file_path)

    print("=" * 60)
    print(f"FILE: {file_path.name}")
    print("=" * 60)

    workbook = load_workbook(file_path, data_only=False)

    # 1. Workbook metadata
    print("\n[METADATA]")
    print("Creator:", workbook.properties.creator)
    print("Last modified by:", workbook.properties.lastModifiedBy)
    print("Title:", workbook.properties.title)
    print("Subject:", workbook.properties.subject)
    print("Company:", getattr(workbook.properties, "company", None))

    # 2. Sheets
    print("\n[SHEETS]")

    for sheet in workbook.worksheets:
        print(
            f"\nSheet: {sheet.title} "
            f"| State: {sheet.sheet_state}"
        )

        # 3. Cell values
        for row in sheet.iter_rows():
            values = []

            for cell in row:
                if cell.value is not None:
                    values.append(
                        f"{cell.coordinate}={cell.value}"
                    )

            if values:
                print(" | ".join(values))


if __name__ == "__main__":
    file_path = input("Enter XLSX file path: ").strip()
    inspect_xlsx(file_path)